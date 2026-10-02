"""Run 享樂遊 (new wording) end to end on the real model, in two steps around the owner's review.

Usage:
  python scripts/run_hsiang_le_you.py import [--serve]
  python scripts/run_hsiang_le_you.py judge [--serve]

import: imports 享樂遊 from the Cathay Century new-wording PDF into
workbooks/享樂遊.new.xlsx, with its Clauses in store/, and imports it a second
time into a scratch directory to check that both give identical extracted
fields. The owner then reviews and confirms the workbook in Excel.

judge: loads the confirmed workbook, recording its extracted and changed field
counts (the first extraction-accuracy figure), then judges each of the
README's two example Scenarios three times, checking that the outcomes are
identical across runs and are the README's Verdicts.

Both steps drive the application interface with the real adapter, configured
by .env as the command line is. They watch the router throughout and record the
most models it held at once, and record the network interfaces present: run
inside `unshare -rn` (loopback only) with --serve, which starts the router in
the same namespace and stops it at the end, to show the run needs no network.
Each step writes results/real-run/<time>-<step>.json, and with --serve the
router's log beside it, and exits non-zero unless every check holds.
"""

import argparse
import contextlib
import hashlib
import json
import os
import signal
import socket
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
from collections.abc import Iterator
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Self

import openpyxl

from travel_claims.app import import_product, judge_scenario, list_policies, load_workbook
from travel_claims.clause_store import ClauseStore
from travel_claims.judging import Reason, Verdict, VerdictMatrix
from travel_claims.llama_server import LlamaServerModels, configured_models
from travel_claims.policies import Wording
from travel_claims.workbook import EXTRACTED

ROOT = Path(__file__).resolve().parent.parent
RESULTS = ROOT / "results" / "real-run"
DOCUMENT = ROOT / "data" / "clauses" / "cathay" / "travel-bundle.new-wording.pdf"
POLICY_NAME = "國泰產物享樂遊海外旅行綜合保險"
PRODUCT = "享樂遊"
WORKBOOK = ROOT / "workbooks" / f"{PRODUCT}.new.xlsx"
STORE = ROOT / "store"
RUNS = 3
# The README's example Scenarios, and the Verdicts it gives for the new wording.
TYPHOON = "投保時海上颱風警報已發布。保險期間內，去程班機因颱風延誤 6 小時。"
ROAD_CLOSURE = (
    "保險期間內，去程班機原定 10:00 起飛，延誤後航空公司安排了 15:00 出發的第一班替代班機。"
    "機場聯外道路封閉，我沒趕上，改搭 20:00 的下一班。"
)
# Router statuses of a model whose instance holds GPU memory.
RESIDENT = {"loading", "loaded", "sleeping"}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("step", choices=["import", "judge"])
    parser.add_argument(
        "--serve", action="store_true", help="start the router for the run and stop it after"
    )
    parser.add_argument(
        "--workbook", type=Path, default=WORKBOOK, help=f"default: {WORKBOOK.relative_to(ROOT)}"
    )
    parser.add_argument("--store", type=Path, default=STORE, help="default: store")
    args = parser.parse_args()
    workbook, store = args.workbook.resolve(), args.store.resolve()
    os.chdir(ROOT)
    models = configured_models()

    started = datetime.now().astimezone()
    results: dict[str, Any] = {
        "step": args.step,
        "started": started.isoformat(timespec="seconds"),
        "model": models.model,
        # The network interfaces of this process's namespace: only lo inside `unshare -rn`.
        "network_interfaces": sorted(name for _, name in socket.if_nameindex()),
    }
    stem = RESULTS / f"{started.strftime('%Y-%m-%dT%H%M%S')}-{args.step}"
    with _router(args.serve, models.base, stem.with_suffix(".log")), Watch(models.base) as watch:
        try:
            if args.step == "import":
                _import(results, workbook, store, models)
            else:
                _judge(results, workbook, store, models)
        except Exception as error:
            results["error"] = f"{type(error).__name__}: {error}"
            raise
        finally:
            results["most_models_resident"] = watch.most
            results["passed"] = (
                "error" not in results
                and all(check for check in results.get("checks", {}).values())
                and watch.most <= 1
            )
            RESULTS.mkdir(parents=True, exist_ok=True)
            path = stem.with_suffix(".json")
            path.write_text(json.dumps(results, ensure_ascii=False, indent=2) + "\n")
            print(json.dumps(results, ensure_ascii=False, indent=2))
            print(f"\n{'PASS' if results['passed'] else 'FAIL'}: {path.relative_to(ROOT)}")
    sys.exit(0 if results["passed"] else 1)


def _import(
    results: dict[str, Any], workbook: Path, store: Path, models: LlamaServerModels
) -> None:
    if workbook.exists():
        sys.exit(f"{workbook} already exists; move it away to import afresh.")
    policies = list_policies(DOCUMENT)
    number = next(n for n, p in enumerate(policies, start=1) if p.name == POLICY_NAME)
    first = _timed_import(number, workbook, ClauseStore(store), models)
    with tempfile.TemporaryDirectory() as scratch:
        again = Path(scratch) / workbook.name
        second = _timed_import(number, again, ClauseStore(Path(scratch) / "store"), models)
        identical = _extracted(workbook) == _extracted(again)
    results |= {
        "document": str(DOCUMENT.relative_to(ROOT)),
        "policy": number,
        "workbook": _shown(workbook),
        "imports": [first, second],
        "extracted_rows_sha256": hashlib.sha256(
            json.dumps(_extracted(workbook), ensure_ascii=False, default=str).encode()
        ).hexdigest(),
        "checks": {"two_imports_identical": identical},
    }


def _timed_import(
    number: int, workbook: Path, store: ClauseStore, models: LlamaServerModels
) -> dict[str, Any]:
    start = time.monotonic()
    imported = import_product(
        DOCUMENT,
        product=PRODUCT,
        wording=Wording.NEW,
        workbook=workbook,
        store=store,
        models=models,
        policy=number,
    )
    return {
        "seconds": round(time.monotonic() - start, 1),
        "clauses": len(imported.policy.clauses),
        "conditions": imported.condition_count,
        "exclusions": imported.exclusion_count,
    }


def _extracted(workbook: Path) -> list[tuple[Any, ...]]:
    """The values as extracted, from the hidden sheet the reviewer does not edit."""
    book = openpyxl.load_workbook(workbook)
    return [tuple(row) for row in book[EXTRACTED].iter_rows(values_only=True)]


def _judge(
    results: dict[str, Any], workbook: Path, directory: Path, models: LlamaServerModels
) -> None:
    store = ClauseStore(directory)
    loaded = load_workbook(workbook, store=store)
    results["workbook"] = _shown(workbook)
    results["load"] = {
        "problems": [str(p) for p in loaded.problems],
        "extracted_fields": loaded.extracted_fields,
        "changed_fields": loaded.changed_fields,
        "unchanged_share": (
            round(1 - loaded.changed_fields / loaded.extracted_fields, 3)
            if loaded.extracted_fields
            else None
        ),
    }
    if loaded.table is None:
        raise RuntimeError(f"{_shown(workbook)} has problems; confirm it first")
    results["confirmed_by"] = loaded.table.confirmed_by
    results["confirmed_on"] = loaded.table.confirmed_on.isoformat()

    scenarios: dict[str, list[dict[str, Any]]] = {}
    for name, scenario in (("typhoon", TYPHOON), ("road closure", ROAD_CLOSURE)):
        runs = []
        for _ in range(RUNS):
            start = time.monotonic()
            matrix = judge_scenario(scenario, [workbook], store=store, models=models)
            runs.append(_summary(matrix) | {"seconds": round(time.monotonic() - start, 1)})
        scenarios[name] = runs
    results["scenarios"] = scenarios

    def outcomes(name: str) -> list[Any]:
        return [run["outcomes"] for run in scenarios[name]]

    typhoon, road = scenarios["typhoon"][0], scenarios["road closure"][0]
    road_readings = {
        reading["verdict"]
        for outcome in road["outcomes"]
        for provision in outcome["turns_on"]
        for reading in provision["readings"]
    }
    results["checks"] = {
        "typhoon_same_in_every_run": all(o == outcomes("typhoon")[0] for o in outcomes("typhoon")),
        "road_closure_same_in_every_run": all(
            o == outcomes("road closure")[0] for o in outcomes("road closure")
        ),
        "typhoon_not_paid_under_31_2": typhoon["verdict"] == Verdict.NOT_PAID
        and any(o["clause"] == "第三十一條 二" for o in typhoon["outcomes"]),
        "road_closure_cause_ambiguous_on_31_5": road["verdict"] == Verdict.UNDETERMINED
        and road["reason"] == Reason.CAUSE_AMBIGUOUS
        and any(o["clause"] == "第三十一條 五" for o in road["outcomes"])
        and road_readings == {Verdict.PAID, Verdict.NOT_PAID},
    }


def _shown(path: Path) -> str:
    """A path relative to the repository when it is inside it."""
    return str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path)


def _summary(matrix: VerdictMatrix) -> dict[str, Any]:
    """The one cell's Verdict and the per-Condition outcomes, without the Clause text."""
    (cell,) = matrix.cells
    return {
        "verdict": str(cell.verdict),
        "reason": None if cell.reason is None else str(cell.reason),
        "outcomes": [
            {
                "condition": o.condition,
                "incident": o.incident,
                "verdict": str(o.verdict),
                "reason": None if o.reason is None else str(o.reason),
                "clause": str(o.clause),
                "delay_minutes": None if o.delay is None else o.delay // timedelta(minutes=1),
                "steps": o.steps,
                "turns_on": [
                    {
                        "provision": f"the {p.kind} of {p.clause}",
                        "readings": [
                            {"cause": r.cause, "verdict": str(r.verdict), "clause": str(r.clause)}
                            for r in p.readings
                        ],
                    }
                    for p in o.turns_on
                ],
            }
            for o in cell.outcomes
        ],
    }


@contextlib.contextmanager
def _router(serve: bool, base: str, log_path: Path) -> Iterator[None]:
    """The router for the run: started here with --serve, else already running."""
    if not serve:
        try:
            _get(base, "/models")
        except OSError:
            sys.exit(f"No router at {base}: start scripts/serve_models.sh, or pass --serve.")
        yield
        return
    with contextlib.suppress(OSError):
        _get(base, "/models")
        sys.exit(f"Something already listens at {base}: stop it, or run without --serve.")
    RESULTS.mkdir(parents=True, exist_ok=True)
    with log_path.open("w") as log:
        server = subprocess.Popen(
            [ROOT / "scripts" / "serve_models.sh"],
            stdout=log,
            stderr=subprocess.STDOUT,
            start_new_session=True,
        )
        try:
            deadline = time.monotonic() + 60
            while True:
                try:
                    _get(base, "/models")
                    break
                except OSError:
                    if server.poll() is not None or time.monotonic() > deadline:
                        sys.exit(f"The router did not come up; see {log_path.relative_to(ROOT)}.")
                    time.sleep(0.5)
            yield
        finally:
            with contextlib.suppress(ProcessLookupError):
                os.killpg(server.pid, signal.SIGTERM)
            server.wait(timeout=60)


class Watch:
    """Polls the router while the run goes on, keeping the most models resident at once."""

    def __init__(self, base: str) -> None:
        self.base = base
        self.most = 0
        self.done = threading.Event()
        self.thread = threading.Thread(target=self._poll, daemon=True)

    def __enter__(self) -> Self:
        self.thread.start()
        return self

    def __exit__(self, *_: object) -> None:
        self.done.set()
        self.thread.join()

    def _poll(self) -> None:
        while not self.done.is_set():
            with contextlib.suppress(OSError, ValueError, KeyError):
                statuses = _get(self.base, "/models")["data"]
                resident = sum(1 for m in statuses if m["status"]["value"] in RESIDENT)
                self.most = max(self.most, resident)
            self.done.wait(0.2)


def _get(base: str, path: str) -> Any:
    with urllib.request.urlopen(base + path, timeout=10) as response:
        return json.load(response)


if __name__ == "__main__":
    main()
