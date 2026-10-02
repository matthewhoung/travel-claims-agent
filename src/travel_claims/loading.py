"""Load a workbook the reviewer has confirmed: validate all of it and list every problem at once.

Nothing is answered from a workbook with a problem, and a missing
confirmation record is one (ADR 0001). Load also counts the extracted fields
and the fields the reviewer changed, which measures extraction accuracy.
"""

import re
from collections import Counter
from collections.abc import Collection
from dataclasses import dataclass, replace
from datetime import date, datetime
from enum import StrEnum
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.worksheet.worksheet import Worksheet

from travel_claims.clause_store import ClauseStore
from travel_claims.conditions import (
    ALL,
    COVERAGE_REQUIREMENTS,
    EXCLUSION_TYPES,
    Amount,
    Availability,
    Benefit,
    BenefitType,
    ClauseRef,
    Condition,
    DelayPeriodRule,
    Exclusion,
    ExclusionType,
)
from travel_claims.policies import Wording
from travel_claims.workbook import (
    AMOUNT_COLUMNS,
    AMOUNTS,
    CONDITION_COLUMNS,
    CONDITIONS,
    CONFIRMED_BY,
    CONFIRMED_ON,
    EXCLUSION_COLUMNS,
    EXCLUSIONS,
    EXTRACTED,
    SHEETS,
    Row,
    YesNo,
    read_confirmation,
    read_extracted,
    read_table,
)

# The reviewer may separate several values in a cell with any of these.
_SEPARATORS = re.compile(r"[;；\n]")
# A date typed as text, such as 2026-10-01 or 2026/10/1.
_DATE = re.compile(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})")
# Chosen at import, not extracted, so not counted as extracted fields.
_CHOSEN_AT_IMPORT = ("Product", "Wording version")
# A reviewed row is matched to an extracted row by the first of these columns
# whose value they share.
_CONDITION_IDENTITY = ("Condition key", "Covered event")
_EXCLUSION_IDENTITY = ("Clause reference", "Text")


@dataclass(frozen=True)
class Problem:
    sheet: str
    # Where it is, such as B1, if it is in one cell.
    cell: str | None
    message: str

    def __str__(self) -> str:
        where = f"{self.sheet}!{self.cell}" if self.cell else self.sheet
        return f"{where}: {self.message}"


@dataclass(frozen=True)
class ConditionTable:
    """A confirmed workbook: the Conditions and exclusions of one Product in one Wording version."""

    product: str
    wording: Wording
    confirmed_by: str
    confirmed_on: date
    conditions: tuple[Condition, ...]
    exclusions: tuple[Exclusion, ...]
    # The rows of the Amounts sheet, in its order.
    amounts: tuple[Amount, ...] = ()

    def amounts_of(self, condition: Condition) -> tuple[Amount, ...]:
        """The Amounts rows of a Condition: those naming its key, or failing that those
        naming all. A Condition with no row is not collected."""
        own = tuple(a for a in self.amounts if a.condition == condition.key)
        every = tuple(
            replace(a, condition=condition.key) for a in self.amounts if a.condition == ALL
        )
        return own or every or (Amount(condition.key, Availability.NOT_COLLECTED),)


@dataclass(frozen=True)
class Loaded:
    # None when there is any problem.
    table: ConditionTable | None
    problems: tuple[Problem, ...]
    # Every field of the rows extracted at import, and how many of them the
    # reviewer changed or deleted.
    extracted_fields: int
    changed_fields: int


def load(path: Path, store: ClauseStore) -> Loaded:
    """Validate the workbook at `path` against the Clauses in `store`, and count changes."""
    book = openpyxl.load_workbook(path, data_only=True)
    missing = [
        Problem(name, None, "the sheet is missing")
        for name in SHEETS
        if name not in book.sheetnames
    ]
    extracted = read_extracted(book[EXTRACTED]) if EXTRACTED in book.sheetnames else None
    if extracted is None and not missing:
        missing.append(Problem(EXTRACTED, None, "the values as extracted are not there"))
    if missing or extracted is None:
        return Loaded(None, tuple(missing), 0, 0)
    product = _StoredProduct(
        extracted.product,
        extracted.wording,
        {clause.number for clause in store.clauses(extracted.product, extracted.wording)},
    )

    problems: list[Problem] = []
    confirmation = _confirmation(book[CONDITIONS], problems)
    tables = {
        name: read_table(book[name], columns)
        for name, columns in (
            (CONDITIONS, CONDITION_COLUMNS),
            (EXCLUSIONS, EXCLUSION_COLUMNS),
            (AMOUNTS, AMOUNT_COLUMNS),
        )
    }
    for name, table in tables.items():
        problems += [
            Problem(name, None, f"the {heading} column is missing")
            for heading in table.missing_columns
        ]
    condition_rows, exclusion_rows = tables[CONDITIONS].rows, tables[EXCLUSIONS].rows
    conditions = _conditions(condition_rows, product, problems)
    targets = _targets(condition_rows)
    exclusions = [
        exclusion
        for row in exclusion_rows
        if (exclusion := _exclusion(row, product, targets, problems)) is not None
    ]
    _force_majeure_readings(condition_rows, conditions, exclusions, problems)
    amounts = _amounts(tables[AMOUNTS].rows, condition_rows, product, problems)

    extracted_fields, changed_fields = _changes(
        (extracted.conditions, condition_rows, CONDITION_COLUMNS, _CONDITION_IDENTITY),
        (extracted.exclusions, exclusion_rows, EXCLUSION_COLUMNS, _EXCLUSION_IDENTITY),
    )
    confirmed = None
    if not problems and confirmation is not None:
        confirmed = ConditionTable(
            product.name,
            product.wording,
            *confirmation,
            tuple(conditions),
            tuple(exclusions),
            tuple(amounts),
        )
    return Loaded(
        confirmed, tuple(sorted(problems, key=_reading_order)), extracted_fields, changed_fields
    )


@dataclass(frozen=True)
class _StoredProduct:
    """The Product and Wording version a workbook was imported as, and its stored Clause numbers."""

    name: str
    wording: Wording
    clauses: Collection[int]

    def __str__(self) -> str:
        return f"{self.name} ({self.wording} wording)"


def _confirmation(sheet: Worksheet, problems: list[Problem]) -> tuple[str, date] | None:
    """Who confirmed the workbook and when, recorded above the Conditions table."""
    record = read_confirmation(sheet)
    found: dict[str, object] = {}
    for label in (CONFIRMED_BY, CONFIRMED_ON):
        if label not in record:
            problems.append(Problem(CONDITIONS, None, f"no {label} above the table"))
            continue
        cell, value = record[label]
        if not _text(value):
            problems.append(Problem(CONDITIONS, cell, f"{label} is missing"))
        elif label == CONFIRMED_ON and _date(value) is None:
            problems.append(Problem(CONDITIONS, cell, f"{label} is not a date: {value}"))
        else:
            found[label] = value
    if len(found) < 2:
        return None
    confirmed_on = _date(found[CONFIRMED_ON])
    assert confirmed_on is not None
    return _text(found[CONFIRMED_BY]) or "", confirmed_on


def _conditions(
    rows: list[Row], product: _StoredProduct, problems: list[Problem]
) -> list[Condition]:
    conditions = []
    first_row: dict[str, int] = {}
    per_benefit = Counter(_text(row.values.get("Benefit")) for row in rows)
    for row in rows:
        fields = _Fields(row, product, problems)
        key = fields.text("Condition key")
        if key in first_row:
            fields.problem("Condition key", f"Condition key {key} is also in row {first_row[key]}")
        first_row.setdefault(key, row.number)
        fields.belongs_to_product()
        benefit = fields.choice("Benefit", Benefit)
        # "Applies to" reads a Benefit's name as all of its Conditions.
        if key in set(Benefit) and key != benefit:
            fields.problem("Condition key", f"Condition key {key} is the name of another Benefit")
        elif key == benefit and per_benefit[key] > 1:
            fields.problem(
                "Condition key",
                f"Condition key {key} is also the name of its Benefit, "
                "which has several Conditions",
            )
        benefit_type = fields.choice("Benefit type", BenefitType)
        flight_delay = benefit is Benefit.FLIGHT_DELAY
        condition = Condition(
            key=key,
            product=product.name,
            wording=product.wording,
            benefit=benefit,
            covered_event=fields.text("Covered event"),
            coverage_requirements=fields.items(
                "Coverage requirements", COVERAGE_REQUIREMENTS.get(benefit)
            ),
            coverage_window=fields.optional_text("Coverage window"),
            threshold_hours=fields.number("Threshold (hours)", required=flight_delay),
            delay_period_rule=fields.optional_choice(
                "Delay-period rule", DelayPeriodRule, required=flight_delay
            ),
            benefit_type=benefit_type,
            step_hours=fields.number(
                "Step (hours)", required=benefit_type is BenefitType.PROGRESSIVE
            ),
            max_claims_per_period=fields.whole("Maximum claims per period"),
            aggregate_limit_group=fields.optional_text("Aggregate limit group"),
            eligible_costs=fields.items("Eligible costs"),
            cost_maximums=fields.optional_text("Cost maximums"),
            clause=fields.clause("Clause reference"),
        )
        if fields.valid:
            conditions.append(condition)
    return conditions


def _targets(rows: list[Row]) -> dict[str, set[Benefit]]:
    """What "Applies to" may name, each with the Benefits it covers.

    A Condition key, a Benefit (all its Conditions), or ALL (every Condition).
    """
    targets: dict[str, set[Benefit]] = {benefit: {benefit} for benefit in Benefit}
    targets[ALL] = set()
    for row in rows:
        key, name = _text(row.values.get("Condition key")), _text(row.values.get("Benefit"))
        benefit = {Benefit(name)} if name in set(Benefit) else set()
        if key is not None:
            targets[key] = benefit
        targets[ALL] |= benefit
    return targets


def _exclusion(
    row: Row, product: _StoredProduct, targets: dict[str, set[Benefit]], problems: list[Problem]
) -> Exclusion | None:
    fields = _Fields(row, product, problems)
    fields.belongs_to_product()
    applies_to = fields.items("Applies to")
    if not applies_to:
        fields.problem("Applies to", "Applies to is missing")
    for name in applies_to:
        if name not in targets:
            fields.problem("Applies to", f"unknown Condition key: {name}")
    exclusion_type = fields.choice("Exclusion type", ExclusionType)
    benefits = {benefit for name in applies_to for benefit in targets.get(name, ())}
    outside = [
        b for b in Benefit if b in benefits and exclusion_type not in EXCLUSION_TYPES.get(b, ())
    ]
    if outside and exclusion_type is not ExclusionType.OTHER:
        fields.problem(
            "Exclusion type",
            f"Exclusion type {exclusion_type} is not on the list for {', '.join(outside)}",
        )
    exclusion = Exclusion(
        product=product.name,
        wording=product.wording,
        type=exclusion_type,
        applies_to=applies_to,
        text=fields.text("Text"),
        proviso=fields.optional_text("Proviso"),
        concerns_cause=fields.choice("Concerns Cause", YesNo) is YesNo.YES,
        clause=fields.clause("Clause reference"),
    )
    return exclusion if fields.valid else None


def _force_majeure_readings(
    rows: list[Row],
    conditions: list[Condition],
    exclusions: list[Exclusion],
    problems: list[Problem],
) -> None:
    """The old-wording rule runs the delay to the next replacement flight when force
    majeure prevented taking the first, as the first-replacement exclusion's proviso
    says; judging takes that reading from the proviso's judgement."""
    row_of = {_text(row.values.get("Condition key")): row for row in rows}
    for condition in conditions:
        if condition.delay_period_rule is not DelayPeriodRule.OLD:
            continue
        if not any(
            e.type is ExclusionType.FIRST_REPLACEMENT_NOT_TAKEN
            and e.proviso is not None
            and e.applies_to_condition(condition)
            for e in exclusions
        ):
            problems.append(
                Problem(
                    CONDITIONS,
                    row_of[condition.key].cell("Delay-period rule"),
                    f"the {DelayPeriodRule.OLD} takes its force-majeure reading from the "
                    f"proviso of a {ExclusionType.FIRST_REPLACEMENT_NOT_TAKEN} exclusion, and "
                    f"no such exclusion with a proviso applies to {condition.key}",
                )
            )


def _amounts(
    rows: list[Row], condition_rows: list[Row], product: _StoredProduct, problems: list[Problem]
) -> list[Amount]:
    """Amounts rows, each naming a Condition, or all of them if it is not published.

    A published amount needs its Plan, its Benefit amount and its Source; the
    others may leave them empty. Amounts are whole NT$.
    """
    keys = {_text(row.values.get("Condition key")) for row in condition_rows} | {ALL}
    amounts = []
    first_row: dict[tuple[str, str | None], int] = {}
    for row in rows:
        fields = _Fields(row, product, problems)
        key = fields.text("Condition key")
        if key and key not in keys:
            fields.problem("Condition key", f"unknown Condition key: {key}")
        availability = fields.optional_choice("Availability", Availability, required=True)
        plan = fields.optional_text("Plan")
        if key and (key, plan) in first_row:
            under = f"under {plan}" if plan else "with no Plan"
            fields.problem("Plan", f"{key} {under} is also in row {first_row[key, plan]}")
        first_row.setdefault((key, plan), row.number)
        amount = Amount(
            condition=key,
            availability=availability or Availability.NOT_COLLECTED,
            plan=plan,
            benefit_amount=fields.nt_dollars("Benefit amount"),
            max_per_incident=fields.nt_dollars("Maximum per incident"),
            source=fields.optional_text("Source"),
        )
        if availability is Availability.PUBLISHED:
            if key == ALL:
                fields.problem(
                    "Condition key", f"a published amount needs a Condition key, not {ALL}"
                )
            for heading in ("Plan", "Benefit amount", "Source"):
                if _text(row.values.get(heading)) is None:
                    fields.problem(heading, f"a published amount needs its {heading}")
        if fields.valid:
            amounts.append(amount)
    return amounts


class _Fields:
    """Reads the fields of one row, recording a problem for each missing or invalid one.

    A required field that is missing or invalid reads as a placeholder, and the
    row is not valid.
    """

    def __init__(self, row: Row, product: _StoredProduct, problems: list[Problem]) -> None:
        self._row = row
        self._product = product
        self._problems = problems
        self.valid = True

    def problem(self, heading: str, message: str) -> None:
        self._problems.append(Problem(self._row.sheet, self._row.cell(heading), message))
        self.valid = False

    def belongs_to_product(self) -> None:
        """Check that the row names the Product and Wording version the workbook was imported as."""
        for heading, expected in (
            ("Product", self._product.name),
            ("Wording version", self._product.wording),
        ):
            value = self.text(heading)
            if value and value != expected:
                self.problem(
                    heading, f"{heading} is {value}, but the workbook is for {self._product}"
                )

    def optional_text(self, heading: str) -> str | None:
        return _text(self._row.values.get(heading))

    def text(self, heading: str) -> str:
        text = self.optional_text(heading)
        if text is None:
            self.problem(heading, f"{heading} is missing")
        return text or ""

    def optional_choice[E: StrEnum](
        self, heading: str, kind: type[E], *, required: bool = False
    ) -> E | None:
        text = self.text(heading) if required else self.optional_text(heading)
        if not text:
            return None
        try:
            return kind(text)
        except ValueError:
            self.problem(heading, f"{heading} is not one of: {', '.join(kind)}")
            return None

    def choice[E: StrEnum](self, heading: str, kind: type[E]) -> E:
        return self.optional_choice(heading, kind, required=True) or next(iter(kind))

    def items(self, heading: str, allowed: Collection[str] | None = None) -> tuple[str, ...]:
        text = self.optional_text(heading) or ""
        items = tuple(item.strip() for item in _SEPARATORS.split(text) if item.strip())
        for item in items:
            if allowed is not None and item not in allowed:
                self.problem(heading, f"{heading}: {item} is not one of: {', '.join(allowed)}")
        return items

    def number(self, heading: str, *, required: bool = False) -> float | None:
        value = self._row.values.get(heading)
        if _text(value) is None:
            if required:
                self.problem(heading, f"{heading} is missing")
            return None
        try:
            return float(str(value).strip())
        except ValueError:
            self.problem(heading, f"{heading} is not a number: {value}")
            return None

    def whole(self, heading: str) -> int | None:
        number = self.number(heading)
        if number is None:
            return None
        if number != int(number) or number < 0:
            self.problem(heading, f"{heading} is not a whole number: {number:g}")
            return None
        return int(number)

    def nt_dollars(self, heading: str) -> int | None:
        """A whole number of NT$; None if the cell is empty."""
        text = _text(self._row.values.get(heading))
        if text is None:
            return None
        try:
            number = float(text)
        except ValueError:
            number = -1
        if number < 0 or not number.is_integer():
            self.problem(heading, f"{heading} is not a whole number of NT$: {text}")
            return None
        return int(number)

    def clause(self, heading: str) -> ClauseRef:
        text = self.text(heading)
        if not text:
            return ClauseRef(0)
        reference = ClauseRef.parse(text)
        if reference is None:
            self.problem(heading, f"{text} is not a Clause reference, such as 第三十一條 二")
            return ClauseRef(0)
        if reference.number not in self._product.clauses:
            self.problem(
                heading,
                f"{ClauseRef(reference.number)} is not among the stored Clauses of {self._product}",
            )
        return reference


def _reading_order(problem: Problem) -> tuple[int, int, int]:
    """Sheet by sheet, then row by row, then column by column."""
    sheet = SHEETS.index(problem.sheet)
    if problem.cell is None:
        return sheet, 0, 0
    column, row = coordinate_from_string(problem.cell)
    return sheet, row, column_index_from_string(column)


def _changes(
    *sheets: tuple[list[dict[str, object]], list[Row], tuple[str, ...], tuple[str, ...]],
) -> tuple[int, int]:
    """How many fields were extracted, and how many of them the reviewer changed.

    Each extracted row is matched to the reviewed row with the same Condition
    key (Clause reference for an exclusion), or failing that the same covered
    event (verbatim text), so rows sorted in Excel still match. A field of a
    row with no match counts as changed; a row the reviewer added has no
    extracted fields.
    """
    extracted = changed = 0
    for before, rows, columns, identity in sheets:
        counted = [heading for heading in columns if heading not in _CHOSEN_AT_IMPORT]
        extracted += len(before) * len(counted)
        unmatched = [row.values for row in rows]
        pending = list(before)
        for heading in identity:
            still_pending = []
            for row in pending:
                same = (
                    a
                    for a in unmatched
                    if _canonical(a.get(heading)) == _canonical(row.get(heading))
                )
                match = next(same, None)
                if match is None:
                    still_pending.append(row)
                    continue
                unmatched.remove(match)
                changed += sum(
                    1 for c in counted if _canonical(row.get(c)) != _canonical(match.get(c))
                )
            pending = still_pending
        changed += len(pending) * len(counted)
    return extracted, changed


def _canonical(value: object) -> str:
    """A cell value as compared: 4, 4.0 and "4" are the same, as are line endings."""
    return (_text(value) or "").replace("\r\n", "\n")


def _text(value: object) -> str | None:
    """A cell value as text, without surrounding spaces; None if empty. 4.0 reads as 4."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = _DATE.fullmatch(str(value).strip())
    if match is None:
        return None
    try:
        return date(*(int(part) for part in match.groups()))
    except ValueError:
        return None
