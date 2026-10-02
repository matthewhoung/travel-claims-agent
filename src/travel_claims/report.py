"""How results are shown: a Verdict matrix's facts as labelled lines and amounts as
text, and the Verdict matrix and the Alignment table in Excel."""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from travel_claims.alignment import AlignmentTable
from travel_claims.conditions import nt_dollars
from travel_claims.judging import PlanAmount, VerdictMatrix
from travel_claims.local_models import (
    UNDATED,
    UNDATED_DAYS,
    ArrangedBy,
    Incident,
    Replacement,
    ScenarioFacts,
)

_MATRIX = "Verdict matrix"
_FACTS = "Facts"
_BREAKDOWN = "Breakdown"
_BREAKDOWN_COLUMNS = (
    "Product",
    "Wording version",
    "Condition key",
    "Incident",
    "Turns on",
    "Reading",
    "Verdict",
    "Reason",
    "Grounds",
    "Clause",
    "Delay period",
    "Steps",
)
_DELAY_COLUMN = _BREAKDOWN_COLUMNS.index("Delay period") + 1
_AMOUNTS = "Amounts"
_ALIGNMENT = "Alignment table"
_AMOUNT_COLUMNS = (
    "Product",
    "Wording version",
    "Condition key",
    "Incident",
    "Plan",
    "Benefit amount",
    "Steps",
    "Total",
    "Source",
)
_WIDE_COLUMNS = {"Grounds", "As read from the Scenario"}


def describe_facts(facts: ScenarioFacts) -> list[tuple[str | None, str]]:
    """The facts read from a Scenario, one line each, labelled.

    A line with no label details the labelled line above it, such as an incident's times.
    """
    lines: list[tuple[str | None, str]] = [("Benefits", ", ".join(facts.benefits) or "none stated")]
    for number, incident in enumerate(facts.incidents, start=1):
        lines.append((f"Incident {number}", _incident(incident)))
        lines.extend((None, detail) for detail in _incident_details(incident))
    lines.append(("Cause", facts.cause or "not stated"))
    lines.append(("Policy bought", _time(facts.purchased_at)))
    if facts.policy_period is not None:
        period = f"{_time(facts.policy_period[0])} – {_time(facts.policy_period[1])}"
    elif facts.within_policy_period is not None:
        where = "within" if facts.within_policy_period else "outside"
        period = f"dates not stated; the trip is {where} it"
    else:
        period = "not stated"
    lines.append(("Policy period", period))
    lines.append(("In force at purchase", "; ".join(facts.in_force_at_purchase) or "none stated"))
    if facts.earlier_claims is not None:
        lines.append(("Earlier claims in the policy period", str(facts.earlier_claims)))
    return lines


def describe_amount(paid: PlanAmount, steps: int | None) -> str:
    """What a paid outcome pays under one Plan, with its source, on one line."""
    if steps is not None and paid.total is None:
        noun = "step" if steps == 1 else "steps"
        paying = (
            f"{nt_dollars(paid.benefit_amount)} per step, {steps} {noun}; "
            "no maximum per incident is given, so no total"
        )
    elif steps is not None and paid.total is not None:
        noun = "step" if steps == 1 else "steps"
        paying = f"{nt_dollars(paid.total)}, {steps} {noun} of {nt_dollars(paid.benefit_amount)}"
        if paid.total < steps * paid.benefit_amount:
            paying += " capped at the maximum per incident"
    elif paid.total is not None:
        paying = nt_dollars(paid.total)
    else:
        paying = f"limit {nt_dollars(paid.benefit_amount)}"
    return f"{paid.plan}: {paying} (source: {paid.source})"


def write_matrix(path: Path, matrix: VerdictMatrix) -> None:
    """Write the Verdict matrix, the facts read, the per-Condition breakdown, and the
    amount per Plan of each paid outcome.

    Each Cause-ambiguous outcome is followed by a row for each reading of each
    provision it turns on. An existing file is never overwritten.
    """
    if path.exists():
        raise FileExistsError(f"{path} already exists; name a new file for the export")
    book = Workbook()
    verdicts = book.active
    assert verdicts is not None  # a new workbook has one sheet
    verdicts.title = _MATRIX
    verdicts.append(["Scenario", matrix.scenario])
    verdicts["A1"].font = Font(bold=True)
    verdicts["B1"].alignment = Alignment(wrap_text=True, vertical="top")
    verdicts.append([])
    _heading(verdicts, ("Product", "Wording version", "Verdict", "Reason"))
    for cell in matrix.cells:
        verdicts.append([cell.product, str(cell.wording), str(cell.verdict), _text(cell.reason)])
    verdicts.column_dimensions["B"].width = 60

    facts = book.create_sheet(_FACTS)
    _heading(facts, ("Fact", "As read from the Scenario"))
    for label, text in describe_facts(matrix.facts):
        facts.append([label, text])

    breakdown = book.create_sheet(_BREAKDOWN)
    _heading(breakdown, _BREAKDOWN_COLUMNS)
    for cell in matrix.cells:
        where = [cell.product, str(cell.wording)]
        for outcome in cell.outcomes:
            condition = [*where, outcome.condition, outcome.incident]
            _breakdown_row(
                breakdown,
                [
                    *condition,
                    None,
                    None,
                    str(outcome.verdict),
                    _text(outcome.reason),
                    outcome.grounds,
                    str(outcome.clause),
                    outcome.delay,
                    outcome.steps,
                ],
            )
            for provision in outcome.turns_on:
                for reading in provision.readings:
                    _breakdown_row(
                        breakdown,
                        [
                            *condition,
                            f"the {provision.kind} of {provision.clause}",
                            reading.cause,
                            str(reading.verdict),
                            _text(reading.reason),
                            reading.grounds,
                            str(reading.clause),
                            reading.delay,
                        ],
                    )

    amounts = book.create_sheet(_AMOUNTS)
    _heading(amounts, _AMOUNT_COLUMNS)
    for cell in matrix.cells:
        for outcome in cell.outcomes:
            for paid in outcome.amounts:
                amounts.append(
                    [
                        cell.product,
                        str(cell.wording),
                        outcome.condition,
                        outcome.incident,
                        paid.plan,
                        paid.benefit_amount,
                        outcome.steps,
                        paid.total,
                        paid.source,
                    ]
                )

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def write_alignment(path: Path, table: AlignmentTable) -> None:
    """Write the Alignment table: a row per row, and for each column its text and what
    each entry cites, one entry per line. An existing file is never overwritten."""
    if path.exists():
        raise FileExistsError(f"{path} already exists; name a new file for the export")
    book = Workbook()
    sheet = book.active
    assert sheet is not None  # a new workbook has one sheet
    sheet.title = _ALIGNMENT
    columns = ["Benefit", "Row"]
    for column in table.columns:
        columns += [str(column), "Cites"]
    _heading(sheet, tuple(columns))
    for row in table.rows:
        values: list[object] = [str(row.benefit) if row.benefit else "all Benefits", row.label]
        for cell, shown in zip(row.cells, row.shown(), strict=True):
            values += [shown, "\n".join(entry.citation() for entry in cell) or None]
        sheet.append(values)
        for written in sheet[sheet.max_row]:
            written.alignment = Alignment(wrap_text=True, vertical="top")
    for index in range(len(table.columns)):
        sheet.column_dimensions[get_column_letter(3 + 2 * index)].width = 60
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def _breakdown_row(sheet: Worksheet, values: list[object]) -> None:
    """A row of the breakdown, its delay period shown in hours and minutes."""
    sheet.append(values)
    if sheet.cell(sheet.max_row, _DELAY_COLUMN).value is not None:
        sheet.cell(sheet.max_row, _DELAY_COLUMN).number_format = "[h]:mm"


def _heading(sheet: Worksheet, columns: tuple[str, ...]) -> None:
    sheet.append(list(columns))
    row = sheet.max_row
    for column, heading in enumerate(columns, start=1):
        sheet.cell(row, column).font = Font(bold=True)
        width = 60 if heading in _WIDE_COLUMNS else 18
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = f"A{row + 1}"


def _text(value: object | None) -> str | None:
    return None if value is None else str(value)


def _incident(incident: Incident) -> str:
    leg = f"{incident.leg} " if incident.leg else ""
    transport = incident.transport or "transport not stated"
    airport = f" from {incident.airport}" if incident.airport else ""
    return f"{leg}{transport}{airport}"


def _incident_details(incident: Incident) -> list[str]:
    times = [f"scheduled departure {_time(incident.scheduled_departure)}"]
    if incident.actual_departure is not None:
        times.append(f"actual departure {_time(incident.actual_departure)}")
    if incident.cancelled:
        times.append("cancelled")
    if incident.missed_connection:
        times.append("connection missed")
    if incident.stated_delay is not None:
        hours = incident.stated_delay.total_seconds() / 3600
        times.append(f"delay stated as {hours:g} hours")
    return [", ".join(times), *(_replacement(r) for r in incident.replacements)]


def _replacement(replacement: Replacement) -> str:
    parts = [f"replacement departing {_time(replacement.departure)}"]
    if replacement.arranged_by is ArrangedBy.INSURED:
        arranged = "arranged by the insured"
        if replacement.arranged_at is not None:
            arranged += f" at {_time(replacement.arranged_at)}"
        parts.append(arranged)
    elif replacement.arranged_by is ArrangedBy.AIRLINE:
        parts.append("arranged by the airline")
    if replacement.destination:
        home = {True: " (Taiwan)", False: " (not Taiwan)", None: ""}
        parts.append(f"to {replacement.destination}{home[replacement.returns_to_taiwan]}")
    if replacement.taken is not None:
        parts.append("taken" if replacement.taken else "not taken")
    return ", ".join(parts)


def _time(moment: datetime | None) -> str:
    if moment is None:
        return "not stated"
    days = (moment.date() - UNDATED).days
    if 0 <= days <= UNDATED_DAYS:
        # A time given without a date, counted in days from the booked flight's.
        later = {0: "", 1: " the next day"}.get(days, f", {days} days later")
        return f"{moment.strftime('%H:%M')}{later} (no date stated)"
    return moment.strftime("%Y-%m-%d %H:%M")
