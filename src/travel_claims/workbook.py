"""The workbook a reviewer confirms in Excel: what extraction hands to everything after it.

Three review sheets: Conditions and Exclusions, filled by extraction, and
Amounts, filled by a person. The reviewer records who confirmed the workbook
and when above the Conditions table. A hidden, protected sheet keeps the
values as extracted, so Load can count the fields the reviewer changed.
"""

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from travel_claims.conditions import (
    Availability,
    Benefit,
    BenefitType,
    Condition,
    DelayPeriodRule,
    Exclusion,
    ExclusionType,
)
from travel_claims.policies import Wording

CONDITIONS = "Conditions"
EXCLUSIONS = "Exclusions"
AMOUNTS = "Amounts"
EXTRACTED = "Extracted"
SHEETS = (CONDITIONS, EXCLUSIONS, AMOUNTS, EXTRACTED)

CONFIRMED_BY = "Confirmed by"
CONFIRMED_ON = "Confirmed on"
# The confirmation record sits above the Conditions table.
CONDITIONS_HEADER_ROW = 4

CONDITION_COLUMNS = (
    "Condition key",
    "Product",
    "Wording version",
    "Benefit",
    "Covered event",
    "Coverage requirements",
    "Coverage window",
    "Threshold (hours)",
    "Delay-period rule",
    "Benefit type",
    "Step (hours)",
    "Maximum claims per period",
    "Aggregate limit group",
    "Eligible costs",
    "Cost maximums",
    "Clause reference",
)
EXCLUSION_COLUMNS = (
    "Product",
    "Wording version",
    "Exclusion type",
    "Applies to",
    "Text",
    "Proviso",
    "Concerns Cause",
    "Clause reference",
)
AMOUNT_COLUMNS = (
    "Condition key",
    "Plan",
    "Availability",
    "Benefit amount",
    "Maximum per incident",
    "Source",
)

# Several values in one cell, such as coverage requirements or the Condition
# keys an exclusion applies to, are separated by this.
LIST_SEPARATOR = "; "


class YesNo(StrEnum):
    YES = "yes"
    NO = "no"


Value = str | int | float | None

# Values a reviewer picks from a list in Excel.
_CHOICES: dict[str, Sequence[str]] = {
    "Wording version": list(Wording),
    "Benefit": list(Benefit),
    "Delay-period rule": list(DelayPeriodRule),
    "Benefit type": list(BenefitType),
    "Exclusion type": list(ExclusionType),
    "Concerns Cause": list(YesNo),
    "Availability": list(Availability),
}
_LAST_CHOICE_ROW = 1000
# The confirmation record and the tables' headings are looked for in the first
# column of this many rows, in case the reviewer inserted rows above them.
_TOP_ROWS = 20
# The hidden sheet starts with the Product and Wording version chosen at import.
_PRODUCT = "Product"
_WORDING = "Wording version"
_WIDE_COLUMNS = {"Covered event", "Coverage window", "Text", "Proviso"}
_TO_FILL = PatternFill("solid", fgColor="FFF2CC")


@dataclass(frozen=True)
class Row:
    """A row of a review sheet, as the reviewer left it."""

    sheet: str
    number: int
    # By column heading.
    values: dict[str, object]
    columns: Mapping[str, int]

    def cell(self, heading: str) -> str | None:
        """Where the value under `heading` is, such as H5."""
        column = self.columns.get(heading)
        return None if column is None else f"{get_column_letter(column)}{self.number}"


@dataclass(frozen=True)
class Table:
    rows: list[Row]
    missing_columns: list[str]


@dataclass(frozen=True)
class ExtractedSheet:
    """What import wrote in the hidden sheet: the rows as extracted, by column heading."""

    product: str
    wording: Wording
    conditions: list[dict[str, object]]
    exclusions: list[dict[str, object]]


def write_draft(
    path: Path,
    product: str,
    wording: Wording,
    conditions: Sequence[Condition],
    exclusions: Sequence[Exclusion],
) -> None:
    """Write a draft workbook, with no confirmation record and no amounts."""
    condition_rows = [condition_row(c) for c in conditions]
    exclusion_rows = [_exclusion_row(e) for e in exclusions]

    book = Workbook()
    sheet = book.active
    assert sheet is not None  # a new workbook has one sheet
    sheet.title = CONDITIONS
    for row, label in enumerate((CONFIRMED_BY, CONFIRMED_ON), start=1):
        sheet.cell(row, 1, label).font = Font(bold=True)
        sheet.cell(row, 2).fill = _TO_FILL
    sheet.cell(2, 2).number_format = "yyyy-mm-dd"
    _table(sheet, CONDITION_COLUMNS, condition_rows, CONDITIONS_HEADER_ROW)
    _table(book.create_sheet(EXCLUSIONS), EXCLUSION_COLUMNS, exclusion_rows)
    _table(book.create_sheet(AMOUNTS), AMOUNT_COLUMNS, [])

    extracted = book.create_sheet(EXTRACTED)
    extracted.append([_PRODUCT, product])
    extracted.append([_WORDING, str(wording)])
    for name, columns, rows in (
        (CONDITIONS, CONDITION_COLUMNS, condition_rows),
        (EXCLUSIONS, EXCLUSION_COLUMNS, exclusion_rows),
    ):
        extracted.append([])
        extracted.append([name])
        extracted.append(columns)
        for values in rows:
            extracted.append([values[heading] for heading in columns])
    extracted.sheet_state = "hidden"
    extracted.protection.sheet = True

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def read_table(sheet: Worksheet, columns: Sequence[str]) -> Table:
    """The rows of the table on `sheet` whose first column is `columns[0]`.

    Columns are found by their headings, so the reviewer may reorder them.
    Empty rows are skipped.
    """
    header_row = _row_labelled(sheet, columns[0])
    if header_row is None:
        return Table([], list(columns))
    found = {
        str(sheet.cell(header_row, column).value): column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(header_row, column).value in columns
    }
    rows = []
    for number in range(header_row + 1, sheet.max_row + 1):
        values: dict[str, object] = {
            heading: sheet.cell(number, column).value for heading, column in found.items()
        }
        if any(_filled(value) for value in values.values()):
            rows.append(Row(sheet.title, number, values, found))
    return Table(rows, [heading for heading in columns if heading not in found])


def read_confirmation(sheet: Worksheet) -> dict[str, tuple[str, object]]:
    """The confirmation record above the Conditions table: each label's cell and value."""
    header_row = _row_labelled(sheet, CONDITION_COLUMNS[0]) or _TOP_ROWS
    record: dict[str, tuple[str, object]] = {}
    for label in (CONFIRMED_BY, CONFIRMED_ON):
        row = _row_labelled(sheet, label)
        if row is not None and row < header_row:
            record[label] = (f"B{row}", sheet.cell(row, 2).value)
    return record


def read_extracted(sheet: Worksheet) -> ExtractedSheet | None:
    """What import wrote in the hidden sheet; None if it is not there as written."""
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2 or rows[0][0] != _PRODUCT or rows[1][0] != _WORDING:
        return None
    if rows[1][1] not in set(Wording):
        return None
    tables: dict[str, list[dict[str, object]]] = {}
    index = 2
    while index < len(rows):
        name = rows[index][0] if rows[index] else None
        if name in (CONDITIONS, EXCLUSIONS) and index + 1 < len(rows):
            headings = [str(h) for h in rows[index + 1] if h is not None]
            index += 2
            table = tables.setdefault(str(name), [])
            while index < len(rows) and any(_filled(value) for value in rows[index]):
                table.append(dict(zip(headings, rows[index], strict=False)))
                index += 1
        else:
            index += 1
    return ExtractedSheet(
        str(rows[0][1]),
        Wording(str(rows[1][1])),
        tables.get(CONDITIONS, []),
        tables.get(EXCLUSIONS, []),
    )


def _row_labelled(sheet: Worksheet, label: str) -> int | None:
    """The first of the top rows whose first cell is `label`."""
    return next((row for row in range(1, _TOP_ROWS + 1) if sheet.cell(row, 1).value == label), None)


def _filled(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def condition_row(c: Condition) -> dict[str, Value]:
    """A Condition's values as its row of the Conditions sheet shows them, by column heading."""
    return {
        "Condition key": c.key,
        "Product": c.product,
        "Wording version": str(c.wording),
        "Benefit": str(c.benefit),
        "Covered event": c.covered_event,
        "Coverage requirements": _joined(c.coverage_requirements),
        "Coverage window": c.coverage_window,
        "Threshold (hours)": _number(c.threshold_hours),
        "Delay-period rule": None if c.delay_period_rule is None else str(c.delay_period_rule),
        "Benefit type": str(c.benefit_type),
        "Step (hours)": _number(c.step_hours),
        "Maximum claims per period": c.max_claims_per_period,
        "Aggregate limit group": c.aggregate_limit_group,
        "Eligible costs": _joined(c.eligible_costs),
        "Cost maximums": c.cost_maximums,
        "Clause reference": str(c.clause),
    }


def _exclusion_row(e: Exclusion) -> dict[str, Value]:
    return {
        "Product": e.product,
        "Wording version": str(e.wording),
        "Exclusion type": str(e.type),
        "Applies to": _joined(e.applies_to),
        "Text": e.text,
        "Proviso": e.proviso,
        "Concerns Cause": str(YesNo.YES if e.concerns_cause else YesNo.NO),
        "Clause reference": str(e.clause),
    }


def _joined(values: Sequence[str]) -> str | None:
    return LIST_SEPARATOR.join(values) or None


def _number(value: float | None) -> int | float | None:
    return int(value) if value is not None and value == int(value) else value


def _table(
    sheet: Worksheet, columns: Sequence[str], rows: list[dict[str, Value]], header_row: int = 1
) -> None:
    for column, heading in enumerate(columns, start=1):
        sheet.cell(header_row, column, heading).font = Font(bold=True)
    for offset, row in enumerate(rows, start=1):
        for column, heading in enumerate(columns, start=1):
            sheet.cell(header_row + offset, column, row[heading])
    sheet.freeze_panes = f"A{header_row + 1}"
    for column, heading in enumerate(columns, start=1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = 48 if heading in _WIDE_COLUMNS else 20
        if heading in _CHOICES:
            choices = DataValidation(
                type="list", formula1=f'"{",".join(_CHOICES[heading])}"', allow_blank=True
            )
            sheet.add_data_validation(choices)
            choices.add(f"{letter}{header_row + 1}:{letter}{_LAST_CHOICE_ROW}")
