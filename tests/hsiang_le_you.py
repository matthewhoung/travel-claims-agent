"""享樂遊 in the new wording, and in the old, as the tests import it.

Each fixture holds its general provisions and its flight-delay Clauses,
captured from the new- and old-wording bundles; the new one also holds the
baggage-delay, baggage-loss and travel-document Clauses. The local models are
scripted to extract the flight-delay Condition, two of its exclusions and one
general exclusion. The old wording has no typhoon exclusion: its 第三十一條 二
is the strike exclusion, and its first-replacement exclusion is 第三十一條 四.
"""

import csv
import re
from dataclasses import replace
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl

from travel_claims.app import import_product
from travel_claims.clause_store import ClauseStore
from travel_claims.conditions import Benefit, BenefitType, ExclusionType
from travel_claims.local_models import (
    ArrangedBy,
    ExtractedCondition,
    ExtractedExclusion,
    Extraction,
    Incident,
    Judgement,
    Leg,
    Reading,
    Replacement,
    ScenarioFacts,
    Settled,
    TurnsOnCause,
)
from travel_claims.policies import Wording

from scripted_models import ScriptedModels

FIXTURE = Path(__file__).parent / "fixtures" / "cathay-new-general-provisions.txt"
# The old-wording bundle's contents page, 享樂遊 on pages 37 to 54, and three riders.
OLD_FIXTURE = Path(__file__).parent / "fixtures" / "cathay-old-hsiang-le-you.txt"
OLD_PAGES = (37, 54)

FLIGHT_DELAY = ExtractedCondition(
    covered_event="scheduled flight departs 4 hours or more late",
    coverage_requirements=("scheduled flight", "as a passenger"),
    threshold_hours=4,
    benefit_type=BenefitType.PROGRESSIVE,
    step_hours=4,
    max_claims_per_period=2,
)
TYPHOON = ExtractedExclusion(
    type=ExclusionType.TYPHOON_WARNING,
    text="要保人向本公司申請訂立保險契約時，中華民國政府氣象機構已發布海上颱風警報。",
    concerns_cause=False,
    item="二",
)
# The force-majeure proviso of 第三十一條 五.
PROVISO = "但被保險人因不可抗力因素致無法搭乘航空業者所提供之第一班替代交通工具者，不在此限。"
FIRST_REPLACEMENT = ExtractedExclusion(
    type=ExclusionType.FIRST_REPLACEMENT_NOT_TAKEN,
    text="被保險人未搭乘航空業者所提供之第一班替代交通工具。",
    concerns_cause=True,
    proviso=PROVISO,
    item="五",
)
WILFUL_ACT = ExtractedExclusion(
    type=ExclusionType.OTHER, text="被保險人故意行為。", concerns_cause=True, item="二"
)

EXTRACTIONS = {
    4: Extraction(exclusions=(WILFUL_ACT,)),
    5: Extraction(policy_period="the dates and times on the policy schedule"),
    30: Extraction(conditions=(FLIGHT_DELAY,)),
    31: Extraction(exclusions=(TYPHOON, FIRST_REPLACEMENT)),
}

# The new wording's strike exclusion, widened to a right to strike already obtained.
STRIKE = ExtractedExclusion(
    type=ExclusionType.STRIKE,
    text="要保人向本公司申請訂立保險契約時，公共交通工具業者之受僱人或機場之地勤、運務人員"
    "已取得罷工權、已預告罷工期間、已宣布罷工或工運活動、已發生罷工或工運活動。",
    concerns_cause=False,
    item="三",
)

# Baggage delay, in the new wording only: the old-wording fixture stops at 第三十二條.
BAGGAGE_DELAY = ExtractedCondition(
    covered_event="checked baggage not received 6 hours after arrival",
    coverage_requirements=(),
    threshold_hours=6,
    benefit_type=BenefitType.ONE_OFF,
    step_hours=None,
    max_claims_per_period=2,
)
RETURN_HOME = ExtractedExclusion(
    type=ExclusionType.RETURN_HOME,
    text="被保險人於返回居住所之行李延誤。",
    concerns_cause=False,
    item="二",
)

OLD_STRIKE = ExtractedExclusion(
    type=ExclusionType.STRIKE,
    text="要保人或被保險人向本公司申請訂立保險契約時，已宣布或已發生罷工或工運活動。",
    concerns_cause=False,
    item="二",
)
# The old wording's 第三十一條 四 has the same proviso, word for word, as the new 五.
OLD_FIRST_REPLACEMENT = replace(FIRST_REPLACEMENT, item="四")

OLD_EXTRACTIONS = {
    4: Extraction(exclusions=(WILFUL_ACT,)),
    5: Extraction(policy_period="the dates and times on the policy schedule"),
    30: Extraction(conditions=(FLIGHT_DELAY,)),
    31: Extraction(exclusions=(OLD_STRIKE, OLD_FIRST_REPLACEMENT)),
}


def import_draft(
    directory: Path, models: ScriptedModels | None = None, wording: Wording = Wording.NEW
) -> tuple[Path, ClauseStore]:
    """Import 享樂遊 in a Wording version; the draft workbook and the Clause store.

    Both Wording versions go to the same Clause store in `directory`.
    Extraction answers with EXTRACTIONS, or OLD_EXTRACTIONS, unless `models`
    is given. The old wording is chosen from its bundle by page range.
    """
    workbook = directory / f"享樂遊.{wording}.xlsx"
    store = ClauseStore(directory / "store")
    old = wording is Wording.OLD
    import_product(
        OLD_FIXTURE if old else FIXTURE,
        product="享樂遊",
        wording=wording,
        workbook=workbook,
        store=store,
        models=models or ScriptedModels(extractions=OLD_EXTRACTIONS if old else EXTRACTIONS),
        pages=OLD_PAGES if old else None,
    )
    return workbook, store


def review(workbook: Path, edits: dict[str, str | float | date | None]) -> None:
    """Edit cells as the reviewer does in Excel, such as {"Conditions!H5": 6}."""
    book = openpyxl.load_workbook(workbook)
    for reference, value in edits.items():
        sheet, cell = reference.split("!")
        book[sheet][cell] = value
    book.save(workbook)


def confirm(workbook: Path, by: str = "王小明", on: date = date(2026, 10, 1)) -> None:
    """Record who confirmed the workbook and when, above the Conditions table."""
    review(workbook, {"Conditions!B1": by, "Conditions!B2": on})


# As a person may type them, valid or not.
AmountRow = tuple[str | float | None, ...]


def enter_amounts(workbook: Path, rows: list[AmountRow]) -> None:
    """Fill the Amounts sheet as a person does, one row per tuple in the sheet's column
    order: Condition key, Plan, availability, Benefit amount, maximum per incident, source."""
    book = openpyxl.load_workbook(workbook)
    for number, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            book["Amounts"].cell(number, column, value)
    book.save(workbook)


# Published Benefit amounts per Plan, committed with a source URL per row.
AMOUNTS_TABLE = Path(__file__).parents[1] / "data" / "amounts.csv"


def cathay_century_flight_delay(*, maximum: bool = True) -> list[AmountRow]:
    """Cathay Century's published flight-delay amounts, as Amounts rows of the flight-delay
    Condition: NT$6,000 per step and at most NT$12,000 per incident, for each Plan.

    Without `maximum`, the maximum per incident is left empty.
    """
    with AMOUNTS_TABLE.open(encoding="utf-8") as file:
        published = [
            row
            for row in csv.DictReader(file)
            if row["insurer"] == "國泰產險" and row["benefit"] == "班機延誤"
        ]
    assert published, f"no published flight-delay amounts of Cathay Century in {AMOUNTS_TABLE}"
    return [
        (
            "flight delay",
            row["tier"],
            "published",
            int(re.search(r"(\d+)元", row["amount_rule"]).group(1)),  # type: ignore[union-attr]
            int(row["max_per_event"]) if maximum else None,
            row["source_url"],
        )
        for row in published
    ]


# Scenarios ------------------------------------------------------------------------

# Times in the policy are Taiwan time (中原標準時間).
TAIPEI = timezone(timedelta(hours=8))
POLICY_PERIOD = (
    datetime(2026, 7, 10, 0, 0, tzinfo=TAIPEI),
    datetime(2026, 7, 14, 23, 59, tzinfo=TAIPEI),
)

# The covered event and its requirements are met, and no exclusion applies,
# in either Wording version.
NOTHING_APPLIES: dict[str, Judgement] = {
    "第三十條": Settled(met=True),
    "第四條 二": Settled(met=False),
    "第三十一條 二": Settled(met=False),
    "第三十一條 四": Settled(met=False),
    "第三十一條 五": Settled(met=False),
}

SCHEDULED = datetime(2026, 7, 10, 8, 0, tzinfo=TAIPEI)


def delayed_by(delay: timedelta) -> ScenarioFacts:
    """The outbound flight from 桃園 departs `delay` late, within the policy period."""
    return ScenarioFacts(
        benefits=(Benefit.FLIGHT_DELAY,),
        incidents=(
            Incident(
                leg=Leg.OUTBOUND,
                airport="桃園國際機場",
                transport="flight",
                scheduled_departure=SCHEDULED,
                actual_departure=SCHEDULED + delay,
            ),
        ),
        cause="機械故障",
        purchased_at=datetime(2026, 7, 1, 12, 0, tzinfo=TAIPEI),
        policy_period=POLICY_PERIOD,
    )


ROAD_CLOSURE = (
    "保險期間內，去程班機原定 10:00 起飛，延誤後航空公司安排了 15:00 出發的第一班替代班機。"
    "機場聯外道路封閉，我沒趕上，改搭 20:00 的下一班。"
)


def road_closure() -> ScenarioFacts:
    """The README's Scenario: the first replacement flight missed because a road was closed.

    It gives times but no dates, and says the trip is within the policy period.
    """
    scheduled = datetime(2026, 7, 10, 10, 0, tzinfo=TAIPEI)
    return ScenarioFacts(
        benefits=(Benefit.FLIGHT_DELAY,),
        incidents=(
            Incident(
                leg=Leg.OUTBOUND,
                airport=None,
                transport="flight",
                scheduled_departure=scheduled,
                replacements=(
                    Replacement(
                        departure=scheduled + timedelta(hours=5),
                        arranged_by=ArrangedBy.AIRLINE,
                        taken=False,
                    ),
                    Replacement(
                        departure=scheduled + timedelta(hours=10),
                        arranged_by=ArrangedBy.AIRLINE,
                        taken=True,
                    ),
                ),
            ),
        ),
        cause="機場聯外道路封閉",
        within_policy_period=True,
    )


FORCE_MAJEURE = TurnsOnCause(
    (Reading(cause="不可抗力", met=True), Reading(cause="非不可抗力", met=False))
)
