"""The command line: a thin shell over the application interface."""

import socket
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

import openpyxl
import pytest

from travel_claims.cli import main
from travel_claims.conditions import Benefit
from travel_claims.local_models import (
    UNDATED,
    ArrangedBy,
    Incident,
    Leg,
    Reading,
    Replacement,
    ScenarioFacts,
    Settled,
    TurnsOnCause,
)
from travel_claims.policies import Wording

import hsiang_le_you
from scripted_models import ScriptedModels

FIXTURES = Path(__file__).parent / "fixtures"


def test_list_policies_prints_each_policy_with_its_pages_wording_and_clauses(
    capsys: pytest.CaptureFixture[str],
) -> None:
    main(["list-policies", str(FIXTURES / "cathay-new-flight-delay.txt")])

    assert capsys.readouterr().out == (
        "1. 國泰產物享樂遊海外旅行綜合保險\n"
        "   pages 13–16, suggested wording: new\n"
        "   4 Clauses: 第一條, 第三十條–第三十二條\n"
    )


def test_list_policies_says_when_no_title_or_wording_was_found(
    capsys: pytest.CaptureFixture[str],
) -> None:
    main(["list-policies", str(FIXTURES / "fubon-old-bundle-boundary.txt")])

    assert capsys.readouterr().out == (
        "1. (no title found)\n"
        "   page 34, no suggested wording (no flight-delay exclusions)\n"
        "   1 Clause: 第三十八條\n"
        "2. 富邦產物安心個人旅行綜合保險-甲型\n"
        "   page 35, no suggested wording (no flight-delay exclusions)\n"
        "   1 Clause: 第一條\n"
    )


def test_list_policies_labels_each_rider_with_its_parent(
    capsys: pytest.CaptureFixture[str],
) -> None:
    main(["list-policies", str(FIXTURES / "cathay-old-hsiang-le-you.txt")])

    out = capsys.readouterr().out
    assert out.startswith(
        "1. 國泰產物享樂遊海外旅行綜合保險\n"
        "   pages 37–44, suggested wording: old\n"
        "   10 Clauses: 第一條–第五條, 第十八條, 第二十七條, 第三十條–第三十二條\n"
        "2. 國泰產物享樂遊海外旅行綜合保險寵物寄宿延長補償保險金附加條款\n"
        "   rider of 國泰產物享樂遊海外旅行綜合保險\n"
        "   page 55, no suggested wording (no flight-delay exclusions)\n"
        "   1 Clause: 第一條\n"
        "3. 國泰產物傷害保險恐怖主義行為保險限額給付附加條款\n"
        "   rider of no policy named in the document\n"
    )


def test_import_writes_a_draft_workbook_and_says_how_to_confirm_it(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = tmp_path / "享樂遊.new.xlsx"
    store = tmp_path / "store"

    main(
        [
            "import",
            str(hsiang_le_you.FIXTURE),
            "--product=享樂遊",
            "--wording=new",
            f"--workbook={workbook}",
            f"--store={store}",
        ],
        models=ScriptedModels(extractions=hsiang_le_you.EXTRACTIONS),
    )

    assert capsys.readouterr().out == (
        "Imported 國泰產物享樂遊海外旅行綜合保險 (pages 13–17) as 享樂遊, new wording.\n"
        "Stored 22 Clauses. Extracted 1 Condition and 3 exclusions into the draft workbook\n"
        f"{workbook}\n"
        "Review it in Excel, and record who confirmed it and when above the Conditions table.\n"
        "Then check it with:\n"
        f"travel-claims load {workbook} --store={store}\n"
    )
    assert workbook.is_file()


def test_load_lists_every_problem_and_the_fields_changed(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)
    load = ["load", str(workbook), f"--store={store.directory}"]

    with pytest.raises(SystemExit) as exit:
        main(load)

    assert exit.value.code == 1
    assert capsys.readouterr().out == (
        "2 problems:\n"
        "Conditions!B1: Confirmed by is missing\n"
        "Conditions!B2: Confirmed on is missing\n"
        "32 fields extracted, 0 changed by the reviewer.\n"
    )

    book = openpyxl.load_workbook(workbook)
    book["Conditions"]["B1"] = "王小明"
    book["Conditions"]["B2"] = date(2026, 10, 1)
    book["Conditions"]["H5"] = 6
    book.save(workbook)
    main(load)

    assert capsys.readouterr().out == (
        "No problems. 享樂遊, new wording: 1 Condition and 3 exclusions,"
        " confirmed by 王小明 on 2026-10-01.\n"
        "32 fields extracted, 1 changed by the reviewer.\n"
    )


def test_import_points_out_a_new_product_name_beside_those_already_stored(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    store = tmp_path / "store"

    def import_as(product: str, wording: str) -> str:
        main(
            [
                "import",
                str(hsiang_le_you.FIXTURE),
                f"--product={product}",
                f"--wording={wording}",
                f"--workbook={tmp_path / f'{product}.{wording}.xlsx'}",
                f"--store={store}",
            ],
            models=ScriptedModels(),
        )
        return capsys.readouterr().out

    first = import_as("享樂遊", "new")
    misspelt = import_as("享樂游", "old")
    same_product = import_as("享樂遊", "old")

    assert "new Product" not in first
    assert (
        "Note: 享樂游 is a new Product. The Clause store already holds 享樂遊; to import"
        " another Wording version of a Product, give its name exactly.\n"
    ) in misspelt
    assert "new Product" not in same_product


def test_judge_prints_the_facts_read_and_the_verdict_matrix(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)
    hsiang_le_you.confirm(workbook)
    taipei = timezone(timedelta(hours=8))
    facts = ScenarioFacts(
        benefits=(Benefit.FLIGHT_DELAY,),
        incidents=(
            Incident(
                leg=Leg.RETURN,
                airport="成田國際機場",
                transport="flight",
                scheduled_departure=datetime(2026, 7, 14, 20, 0, tzinfo=taipei),
                cancelled=True,
                replacements=(
                    Replacement(
                        departure=datetime(2026, 7, 15, 14, 0, tzinfo=taipei),
                        arranged_by=ArrangedBy.INSURED,
                        taken=True,
                        arranged_at=datetime(2026, 7, 15, 9, 0, tzinfo=taipei),
                        destination="桃園",
                        returns_to_taiwan=True,
                    ),
                ),
            ),
        ),
        cause="颱風",
        purchased_at=datetime(2026, 7, 1, 12, 0, tzinfo=taipei),
        policy_period=(
            datetime(2026, 7, 10, 0, 0, tzinfo=taipei),
            datetime(2026, 7, 14, 23, 59, tzinfo=taipei),
        ),
    )
    models = ScriptedModels(
        facts=facts,
        judgements={
            "第三十條": Settled(met=True),
            "第四條 二": Settled(met=False),
            "第三十一條 二": Settled(met=False),
            "第三十一條 五": Settled(met=False),
        },
    )

    main(
        ["judge", "--scenario=回程班機因颱風取消。", str(workbook), f"--store={store.directory}"],
        models=models,
    )

    assert models.scenarios == ["回程班機因颱風取消。"]
    assert capsys.readouterr().out == (
        "Facts read from the Scenario:\n"
        "  Benefits: flight delay\n"
        "  Incident 1: return flight from 成田國際機場\n"
        "    scheduled departure 2026-07-14 20:00, cancelled\n"
        "    replacement departing 2026-07-15 14:00, arranged by the insured"
        " at 2026-07-15 09:00, to 桃園 (Taiwan), taken\n"
        "  Cause: 颱風\n"
        "  Policy bought: 2026-07-01 12:00\n"
        "  Policy period: 2026-07-10 00:00 – 2026-07-14 23:59\n"
        "  In force at purchase: none stated\n"
        "\n"
        "享樂遊, new wording: paid\n"
        "  flight delay, incident 1: paid (第三十條)\n"
        "    a delay of 18 h 0 min: 4 full steps of 4 hours\n"
    )


def test_judge_prints_the_amount_per_plan_of_a_paid_outcome(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)
    hsiang_le_you.confirm(workbook)
    hsiang_le_you.enter_amounts(
        workbook,
        [
            ("flight delay", "A", "published", 6000, 12000, "https://example.com/a"),
            ("flight delay", "B", "published", 6000, None, "https://example.com/b"),
            ("flight delay", "C", "not published", None, None, None),
        ],
    )
    models = ScriptedModels(
        facts=hsiang_le_you.delayed_by(timedelta(hours=12)),
        judgements=hsiang_le_you.NOTHING_APPLIES,
    )

    main(
        ["judge", "--scenario=延誤十二小時。", str(workbook), f"--store={store.directory}"],
        models=models,
    )

    out = capsys.readouterr().out
    assert out.endswith(
        "  flight delay, incident 1: paid (第三十條)\n"
        "    a delay of 12 h 0 min: 3 full steps of 4 hours\n"
        "    A: NT$12,000, 3 steps of NT$6,000 capped at the maximum per incident"
        " (source: https://example.com/a)\n"
        "    B: NT$6,000 per step, 3 steps; no maximum per incident is given, so no total"
        " (source: https://example.com/b)\n"
    )


def test_judge_prints_each_reading_of_a_cause_ambiguous_outcome_and_exports_the_matrix(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)
    hsiang_le_you.confirm(workbook)
    models = ScriptedModels(
        facts=hsiang_le_you.road_closure(),
        judgements=hsiang_le_you.NOTHING_APPLIES
        | {
            "第三十一條 五": Settled(met=True),
            hsiang_le_you.PROVISO: hsiang_le_you.FORCE_MAJEURE,
        },
    )
    exported = tmp_path / "verdicts.xlsx"

    main(
        [
            "judge",
            f"--scenario={hsiang_le_you.ROAD_CLOSURE}",
            str(workbook),
            f"--store={store.directory}",
            f"--export={exported}",
        ],
        models=models,
    )

    out = capsys.readouterr().out
    assert out.endswith(
        "享樂遊, new wording: undetermined, Cause ambiguous\n"
        "  flight delay, incident 1: undetermined, Cause ambiguous (第三十一條 五)\n"
        f"    turns on how the Cause is classified, under the proviso: {hsiang_le_you.PROVISO}\n"
        "    if the proviso of 第三十一條 五 is read as 不可抗力: paid (第三十條)\n"
        "      a delay of 5 h 0 min: 1 full step of 4 hours\n"
        "    if the proviso of 第三十一條 五 is read as 非不可抗力: not paid (第三十一條 五)\n"
        "      the exclusion applies: 被保險人未搭乘航空業者所提供之第一班替代交通工具。\n"
        "\n"
        f"Exported the Verdict matrix to {exported}\n"
    )
    assert openpyxl.load_workbook(exported).sheetnames == [
        "Verdict matrix",
        "Facts",
        "Breakdown",
        "Amounts",
    ]


def test_judge_refuses_to_export_over_an_existing_file_before_judging(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)
    hsiang_le_you.confirm(workbook)
    models = ScriptedModels(facts=hsiang_le_you.road_closure())

    with pytest.raises(SystemExit) as exit:
        main(
            [
                "judge",
                "--scenario=班機延誤五小時。",
                str(workbook),
                f"--store={store.directory}",
                f"--export={workbook}",
            ],
            models=models,
        )

    assert exit.value.code == 2
    assert f"{workbook} already exists" in capsys.readouterr().err
    assert models.scenarios == []


def test_judge_reports_an_answer_the_local_models_may_not_give(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)
    hsiang_le_you.confirm(workbook)
    # The typhoon-warning exclusion does not concern the Cause.
    models = ScriptedModels(
        facts=hsiang_le_you.road_closure(),
        judgements=hsiang_le_you.NOTHING_APPLIES
        | {"第三十一條 二": TurnsOnCause((Reading("颱風", met=True), Reading("其他", met=False)))},
    )

    with pytest.raises(SystemExit) as exit:
        main(
            [
                "judge",
                f"--scenario={hsiang_le_you.ROAD_CLOSURE}",
                str(workbook),
                f"--store={store.directory}",
            ],
            models=models,
        )

    assert exit.value.code == 1
    assert capsys.readouterr().out == (
        "Nothing judged: the local models answered that the exclusion of 第三十一條 二 turns on "
        "the Cause, which only a provision that concerns the Cause may: "
        "要保人向本公司申請訂立保險契約時，中華民國政府氣象機構已發布海上颱風警報。\n"
    )


def test_judge_refuses_a_workbook_with_problems_and_lists_them(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)

    with pytest.raises(SystemExit) as exit:
        main(
            ["judge", "--scenario=班機延誤五小時。", str(workbook), f"--store={store.directory}"],
            models=ScriptedModels(),
        )

    assert exit.value.code == 1
    assert capsys.readouterr().out == (
        "Nothing judged: fix these problems first.\n"
        f"{workbook}: 2 problems\n"
        "Conditions!B1: Confirmed by is missing\n"
        "Conditions!B2: Confirmed on is missing\n"
    )


def test_import_uses_the_local_models_configured_and_says_when_they_are_not_running(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        closed = probe.getsockname()[1]
    monkeypatch.chdir(tmp_path)
    (tmp_path / ".env").write_text(f"LLAMA_PORT={closed}\nLLM_MODEL=qwen3.5-4b\n")

    with pytest.raises(SystemExit) as exit:
        main(["import", str(hsiang_le_you.FIXTURE), "--product=享樂遊", "--wording=new"])

    assert exit.value.code == 1
    error = capsys.readouterr().err
    assert f"llama-server is not reachable at http://127.0.0.1:{closed}" in error
    assert "scripts/serve_models.sh" in error
    assert not (tmp_path / "workbooks").exists()


def test_judge_prints_times_given_without_a_date_as_times_of_day(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)
    hsiang_le_you.confirm(workbook)
    scheduled = datetime.combine(UNDATED, time(22, 0), hsiang_le_you.TAIPEI)
    facts = ScenarioFacts(
        benefits=(Benefit.FLIGHT_DELAY,),
        incidents=(
            Incident(
                leg=Leg.OUTBOUND,
                airport=None,
                transport="flight",
                scheduled_departure=scheduled,
                replacements=(
                    Replacement(
                        departure=scheduled + timedelta(hours=3),
                        arranged_by=ArrangedBy.AIRLINE,
                        taken=True,
                    ),
                ),
            ),
        ),
        within_policy_period=True,
    )
    models = ScriptedModels(facts=facts, judgements=hsiang_le_you.NOTHING_APPLIES)

    main(
        [
            "judge",
            "--scenario=去程班機原定 22:00 起飛……",
            str(workbook),
            f"--store={store.directory}",
        ],
        models=models,
    )

    facts_read = capsys.readouterr().out.split("\n\n")[0]
    assert "    scheduled departure 22:00 (no date stated)\n" in facts_read
    assert (
        "    replacement departing 01:00 the next day (no date stated), arranged by the airline"
        in facts_read
    )


def test_align_prints_the_alignment_table_and_exports_it(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    old, store = hsiang_le_you.import_draft(tmp_path, wording=Wording.OLD)
    new, _ = hsiang_le_you.import_draft(tmp_path)
    hsiang_le_you.confirm(old)
    hsiang_le_you.confirm(new)
    hsiang_le_you.enter_amounts(new, hsiang_le_you.cathay_century_flight_delay())
    exported = tmp_path / "alignment.xlsx"

    main(["align", str(new), str(old), f"--store={store.directory}", f"--export={exported}"])

    out = capsys.readouterr().out
    assert out.startswith(
        "Columns: 享樂遊 (old wording), 享樂遊 (new wording)\n"
        "\n"
        "flight delay\n"
        "  Covered event\n"
        "    享樂遊 (old wording): scheduled flight departs 4 hours or more late (第三十條)\n"
        "    享樂遊 (new wording): scheduled flight departs 4 hours or more late (第三十條)\n"
    )
    assert (
        "  Benefit amount\n"
        "    享樂遊 (old wording): not available: not collected (第三十條)\n"
        "    享樂遊 (new wording): 安心型(T5): NT$6,000 per step, at most NT$12,000 per incident"
        " (第三十條; source: https://www.cathay-ins.com.tw/cathayins/personal/travel/oversea/)\n"
    ) in out
    assert (
        "  typhoon warning at purchase\n"
        "    享樂遊 (old wording): absent\n"
        "    享樂遊 (new wording): 要保人向本公司申請訂立保險契約時，中華民國政府氣象機構"
        "已發布海上颱風警報。 (第三十一條 二)\n"
    ) in out
    # An "other" exclusion's row shows only the column it is in.
    assert out.endswith(
        "all Benefits\n"
        "  other\n"
        "    享樂遊 (old wording): 被保險人故意行為。 (第四條 二)\n"
        "  other\n"
        "    享樂遊 (new wording): 被保險人故意行為。 (第四條 二)\n"
        "\n"
        f"Exported the Alignment table to {exported}\n"
    )
    assert openpyxl.load_workbook(exported).sheetnames == ["Alignment table"]


def test_align_refuses_a_workbook_with_problems_and_lists_them(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)
    exported = tmp_path / "alignment.xlsx"

    with pytest.raises(SystemExit) as exit:
        main(["align", str(workbook), f"--store={store.directory}", f"--export={exported}"])

    assert exit.value.code == 1
    assert capsys.readouterr().out == (
        "No Alignment table built: fix these problems first.\n"
        f"{workbook}: 2 problems\n"
        "Conditions!B1: Confirmed by is missing\n"
        "Conditions!B2: Confirmed on is missing\n"
    )
    assert not exported.exists()


def test_align_refuses_to_export_over_an_existing_file(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook, store = hsiang_le_you.import_draft(tmp_path)
    hsiang_le_you.confirm(workbook)

    with pytest.raises(SystemExit) as exit:
        main(["align", str(workbook), f"--store={store.directory}", f"--export={workbook}"])

    assert exit.value.code == 2
    assert f"{workbook} already exists" in capsys.readouterr().err
