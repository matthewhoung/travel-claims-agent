"""Import a Product: one policy of a clause document becomes a draft workbook.

The local models are scripted (tests/scripted_models.py): each test says what
extraction returns for which Clause.
"""

from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet

from travel_claims.app import CannotImport, import_product, load_workbook
from travel_claims.clause_store import ClauseStore
from travel_claims.conditions import Benefit, BenefitType
from travel_claims.local_models import ClauseRole, Extraction
from travel_claims.policies import Wording

from hsiang_le_you import (
    BAGGAGE_DELAY,
    CANCELLED_BY_RELATIVES_DEATH,
    CANCELLED_BY_STRIKE,
    FIRST_REPLACEMENT,
    FIXTURE,
    FLIGHT_DELAY,
    RETURN_HOME,
    TRIP_EXTRACTIONS,
    TYPHOON,
    WILFUL_ACT,
    confirm,
    import_draft,
)
from scripted_models import ScriptedModels

# 享樂遊 on pages 13 to 19, then 享暢行 on page 20.
BUNDLE = Path(__file__).parent / "fixtures" / "cathay-new-bundle-openings.txt"


def table(sheet: Worksheet, header_row: int = 1) -> list[dict[str, object]]:
    """The rows of a sheet's table, as a reviewer sees them, keyed by column heading."""
    rows = sheet.iter_rows(min_row=header_row, values_only=True)
    headings = [str(heading) for heading in next(rows)]
    return [dict(zip(headings, row, strict=True)) for row in rows if any(row)]


def headings(sheet: Worksheet, header_row: int = 1) -> list[object]:
    return [cell.value for cell in sheet[header_row]]


def test_import_drafts_a_workbook_with_the_flight_delay_condition(tmp_path: Path) -> None:
    models = ScriptedModels(
        extractions={
            5: Extraction(policy_period="the dates and times on the policy schedule"),
            6: Extraction(policy_period="extended up to 24 hours for a delayed arrival"),
            30: Extraction(conditions=(FLIGHT_DELAY,)),
        }
    )
    book = imported(tmp_path, models)

    assert book.sheetnames == ["Conditions", "Exclusions", "Amounts", "Extracted"]
    assert table(book["Conditions"], header_row=4) == [
        {
            "Condition key": "flight delay",
            "Product": "享樂遊",
            "Wording version": "new",
            "Benefit": "flight delay",
            "Covered event": "scheduled flight departs 4 hours or more late",
            "Coverage requirements": "scheduled flight; as a passenger",
            "Coverage window": "the dates and times on the policy schedule\n"
            "extended up to 24 hours for a delayed arrival",
            "Window (days before departure)": None,
            "Threshold (hours)": 4,
            "Delay-period rule": "new-wording rule",
            "Benefit type": "progressive fixed amount",
            "Step (hours)": 4,
            "Maximum claims per period": 2,
            "Aggregate limit group": None,
            "Eligible costs": None,
            "Cost maximums": None,
            "Clause reference": "第三十條",
        }
    ]
    assert headings(book["Amounts"]) == [
        "Condition key",
        "Plan",
        "Availability",
        "Benefit amount",
        "Maximum per incident",
        "Source",
    ]
    assert table(book["Amounts"]) == []


def test_an_exclusion_applies_to_its_benefits_conditions_or_if_general_to_all(
    tmp_path: Path,
) -> None:
    models = ScriptedModels(
        extractions={
            4: Extraction(exclusions=(WILFUL_ACT,)),
            30: Extraction(conditions=(FLIGHT_DELAY,)),
            31: Extraction(exclusions=(TYPHOON, FIRST_REPLACEMENT)),
        }
    )

    book = imported(tmp_path, models)

    assert table(book["Exclusions"]) == [
        {
            "Product": "享樂遊",
            "Wording version": "new",
            "Exclusion type": "other",
            "Applies to": "all",
            "Text": "被保險人故意行為。",
            "Proviso": None,
            "Concerns Cause": "yes",
            "Clause reference": "第四條 二",
        },
        {
            "Product": "享樂遊",
            "Wording version": "new",
            "Exclusion type": "typhoon warning at purchase",
            "Applies to": "flight delay",
            "Text": "要保人向本公司申請訂立保險契約時，中華民國政府氣象機構已發布海上颱風警報。",
            "Proviso": None,
            "Concerns Cause": "no",
            "Clause reference": "第三十一條 二",
        },
        {
            "Product": "享樂遊",
            "Wording version": "new",
            "Exclusion type": "first replacement not taken",
            "Applies to": "flight delay",
            "Text": "被保險人未搭乘航空業者所提供之第一班替代交通工具。",
            "Proviso": "但被保險人因不可抗力因素致無法搭乘航空業者所提供之第一班替代交通工具者，"
            "不在此限。",
            "Concerns Cause": "yes",
            "Clause reference": "第三十一條 五",
        },
    ]


def test_condition_keys_are_the_benefit_and_a_label_unique_and_stable(tmp_path: Path) -> None:
    models = ScriptedModels(
        extractions={
            30: Extraction(
                conditions=(
                    FLIGHT_DELAY,
                    replace(FLIGHT_DELAY, label="missed connection", item="三"),
                    replace(FLIGHT_DELAY, label="missed connection"),
                )
            ),
            31: Extraction(exclusions=(TYPHOON,)),
        }
    )

    first = imported(tmp_path / "first", models)
    again = imported(tmp_path / "again", models)

    # A key never repeats, and never names the Benefit, which "Applies to"
    # uses for all of its Conditions.
    keys = [
        "flight delay (1)",
        "flight delay / missed connection (1)",
        "flight delay / missed connection (2)",
    ]
    for book in (first, again):
        conditions = table(book["Conditions"], header_row=4)
        assert [row["Condition key"] for row in conditions] == keys
        assert [row["Clause reference"] for row in conditions] == [
            "第三十條",
            "第三十條 三",
            "第三十條",
        ]
        assert [row["Applies to"] for row in table(book["Exclusions"])] == ["; ".join(keys)]


def test_the_general_provisions_and_the_extracted_benefits_are_extracted_and_every_clause_indexed(
    tmp_path: Path,
) -> None:
    models = ScriptedModels()

    imported(tmp_path, models)

    # Trip cancellation, flight delay and trip change are judged; baggage delay,
    # baggage loss and loss of travel documents are extracted for the Alignment
    # table only. Claim documents and recovery Clauses are not extracted.
    assert [(r.clause.number, r.role, r.benefit) for r in models.extracted] == [
        (3, ClauseRole.DEFINITIONS, None),
        (4, ClauseRole.GENERAL_EXCLUSIONS, None),
        (5, ClauseRole.POLICY_PERIOD, None),
        (6, ClauseRole.POLICY_PERIOD, None),
        (27, ClauseRole.BENEFIT_COVER, Benefit.TRIP_CANCELLATION),
        (28, ClauseRole.BENEFIT_EXCLUSIONS, Benefit.TRIP_CANCELLATION),
        (30, ClauseRole.BENEFIT_COVER, Benefit.FLIGHT_DELAY),
        (31, ClauseRole.BENEFIT_EXCLUSIONS, Benefit.FLIGHT_DELAY),
        (33, ClauseRole.BENEFIT_COVER, Benefit.TRIP_CHANGE),
        (34, ClauseRole.BENEFIT_EXCLUSIONS, Benefit.TRIP_CHANGE),
        (36, ClauseRole.BENEFIT_COVER, Benefit.BAGGAGE_DELAY),
        (37, ClauseRole.BENEFIT_EXCLUSIONS, Benefit.BAGGAGE_DELAY),
        (39, ClauseRole.BENEFIT_COVER, Benefit.BAGGAGE_LOSS),
        (40, ClauseRole.BENEFIT_EXCLUSIONS, Benefit.BAGGAGE_LOSS),
        (41, ClauseRole.BENEFIT_EXCLUSIONS, Benefit.BAGGAGE_LOSS),
        (45, ClauseRole.BENEFIT_COVER, Benefit.TRAVEL_DOCUMENT_LOSS),
        (46, ClauseRole.BENEFIT_EXCLUSIONS, Benefit.TRAVEL_DOCUMENT_LOSS),
    ]
    # Each with the other Clauses of its chapter as context.
    context = {r.clause.number: [c.number for c in r.context] for r in models.extracted}
    assert context[4] == [1, 3, 5, 6]
    assert context[30] == [27, 28, 31, 32, 33, 34, *range(36, 48)]
    assert models.indexed == [
        ("享樂遊", Wording.NEW, (1, 3, 4, 5, 6, 18, 27, 28, 30, 31, 32, 33, 34, *range(36, 48)))
    ]


def test_trip_cancellation_and_trip_change_are_drafted_one_condition_per_covered_cause(
    tmp_path: Path,
) -> None:
    book = imported(tmp_path, ScriptedModels(extractions=TRIP_EXTRACTIONS))

    conditions = table(book["Conditions"], header_row=4)
    trips = [c for c in conditions if c["Benefit"] in ("trip cancellation", "trip change")]
    assert [
        (
            c["Condition key"],
            c["Coverage requirements"],
            c["Coverage window"],
            c["Window (days before departure)"],
            c["Benefit type"],
            c["Eligible costs"],
            c["Aggregate limit group"],
            c["Clause reference"],
        )
        for c in trips
    ] == [
        (
            "trip cancellation / relative's death",
            "death or critical illness of the insured or a relative",
            "自預定海外旅程開始前二十日起至海外旅行期間開始時止",
            20,
            "reimbursement",
            "tour fee; transport; lodging; tickets",
            "trip cancellation",
            "第二十七條 一",
        ),
        (
            "trip cancellation / strike",
            "strike cancelling or delaying the booked transport",
            "自預定海外旅程開始前二十日起至海外旅行期間開始時止",
            20,
            "reimbursement",
            "tour fee; transport; lodging; tickets",
            "trip cancellation",
            "第二十七條 三",
        ),
        (
            "trip change / strike",
            "strike of the booked transport",
            "海外旅行期間內",
            None,
            "reimbursement",
            "transport; lodging",
            "trip change",
            "第三十三條 一",
        ),
        (
            "trip change / relative's death",
            "death or critical illness of a spouse or relative in Taiwan",
            "海外旅行期間內",
            None,
            "reimbursement",
            "transport; lodging",
            "trip change",
            "第三十三條 三",
        ),
    ]
    assert trips[2]["Cost maximums"] == (
        "原預定之交通或每日住宿費用各增加20%；無證明者每日合計新臺幣2,000元"
    )
    assert [
        (e["Exclusion type"], e["Applies to"], e["Clause reference"])
        for e in table(book["Exclusions"])
        if str(e["Clause reference"]).startswith(("第二十八條", "第三十四條"))
    ] == [
        (
            "incident already occurred at purchase",
            "trip cancellation / relative's death; trip cancellation / strike",
            "第二十八條 三",
        ),
        (
            "incident already occurred at purchase",
            "trip change / strike; trip change / relative's death",
            "第三十四條 二",
        ),
        (
            "first replacement not taken",
            "trip change / strike; trip change / relative's death",
            "第三十四條 六",
        ),
    ]


def test_conditions_share_an_aggregate_limit_group_only_where_their_clause_caps_the_period_total(
    tmp_path: Path,
) -> None:
    def groups(directory: Path, cover: Extraction) -> list[object]:
        book = imported(directory, ScriptedModels(extractions={27: cover}))
        return [c["Aggregate limit group"] for c in table(book["Conditions"], header_row=4)]

    both = (CANCELLED_BY_RELATIVES_DEATH, CANCELLED_BY_STRIKE)
    # The Clause caps the total paid in the policy period, and the Benefit has
    # several reimbursement Conditions.
    assert groups(tmp_path / "capped", Extraction(both, caps_period_total=True)) == [
        "trip cancellation",
        "trip cancellation",
    ]
    assert groups(tmp_path / "uncapped", Extraction(both)) == [None, None]
    # One Condition shares its limit with no other.
    one = Extraction((CANCELLED_BY_RELATIVES_DEATH,), caps_period_total=True)
    assert groups(tmp_path / "one", one) == [None]
    # Fixed amounts are not reimbursed, so there is no total of reimbursements to cap.
    fixed = tuple(replace(c, benefit_type=BenefitType.ONE_OFF) for c in both)
    assert groups(tmp_path / "fixed", Extraction(fixed, caps_period_total=True)) == [None, None]


def test_an_alignment_only_benefits_exclusions_apply_to_its_own_conditions(
    tmp_path: Path,
) -> None:
    models = ScriptedModels(
        extractions={
            30: Extraction(conditions=(FLIGHT_DELAY,)),
            36: Extraction(conditions=(BAGGAGE_DELAY,)),
            37: Extraction(exclusions=(RETURN_HOME,)),
        }
    )

    book = imported(tmp_path, models)

    assert [
        (row["Condition key"], row["Benefit"], row["Benefit type"], row["Clause reference"])
        for row in table(book["Conditions"], header_row=4)
    ] == [
        ("flight delay", "flight delay", "progressive fixed amount", "第三十條"),
        ("baggage delay", "baggage delay", "one-off fixed amount", "第三十六條"),
    ]
    assert [
        (row["Exclusion type"], row["Applies to"], row["Clause reference"])
        for row in table(book["Exclusions"])
    ] == [("delay on return home", "baggage delay", "第三十七條 二")]


def test_conditions_from_a_clause_that_is_not_a_benefits_cover_are_refused(
    tmp_path: Path,
) -> None:
    # Definitions state no payout; a Condition read from them has no Benefit.
    models = ScriptedModels(extractions={3: Extraction(conditions=(FLIGHT_DELAY,))})

    with pytest.raises(ValueError, match="Conditions for 第三條 用詞定義, which is not"):
        import_draft(tmp_path, models)


def test_the_extracted_values_are_kept_in_a_hidden_protected_sheet(tmp_path: Path) -> None:
    book = imported(
        tmp_path, ScriptedModels(extractions={30: Extraction(conditions=(FLIGHT_DELAY,))})
    )

    extracted = book["Extracted"]
    assert extracted.sheet_state == "hidden"
    assert extracted.protection.sheet


def test_the_policy_is_chosen_by_its_number_in_the_list_or_by_a_page_range(
    tmp_path: Path,
) -> None:
    by_number, by_pages = ScriptedModels(), ScriptedModels()
    store = ClauseStore(tmp_path / "store")

    imported_policy = import_product(
        BUNDLE,
        product="享暢行",
        wording=Wording.NEW,
        workbook=tmp_path / "by-number.xlsx",
        store=store,
        models=by_number,
        policy=2,
    )
    import_product(
        BUNDLE,
        product="享暢行",
        wording=Wording.NEW,
        workbook=tmp_path / "by-pages.xlsx",
        store=store,
        models=by_pages,
        pages=(20, 20),
    )

    assert imported_policy.policy.name == "國泰產物享暢行海外旅行綜合保險"
    assert by_number.indexed == by_pages.indexed == [("享暢行", Wording.NEW, (1, 2))]


@pytest.mark.parametrize(
    ("policy", "pages", "message"),
    [
        (None, None, "the document holds 2 policies; choose one by its number in the list"),
        (3, None, "there is no policy 3; the document holds 2"),
        (None, (21, 22), "pages 21–22 hold no policy"),
        (None, (20, 99), "pages 20–99 are not in the document, which has 27 pages"),
    ],
)
def test_import_refuses_a_choice_that_names_no_single_policy(
    tmp_path: Path, policy: int | None, pages: tuple[int, int] | None, message: str
) -> None:
    with pytest.raises(CannotImport, match=message):
        import_product(
            BUNDLE,
            product="享暢行",
            wording=Wording.NEW,
            workbook=tmp_path / "享暢行.new.xlsx",
            store=ClauseStore(tmp_path / "store"),
            models=ScriptedModels(),
            policy=policy,
            pages=pages,
        )


def test_import_never_overwrites_a_workbook(tmp_path: Path) -> None:
    workbook = tmp_path / "享樂遊.new.xlsx"
    workbook.write_bytes(b"a confirmed workbook")
    models = ScriptedModels()

    with pytest.raises(CannotImport, match="already exists"):
        import_product(
            FIXTURE,
            product="享樂遊",
            wording=Wording.NEW,
            workbook=workbook,
            store=ClauseStore(tmp_path / "store"),
            models=models,
        )

    assert workbook.read_bytes() == b"a confirmed workbook"
    assert not models.extracted
    assert not models.indexed


def test_another_wording_version_of_a_product_is_stored_beside_the_first(tmp_path: Path) -> None:
    new, store = import_draft(tmp_path)
    old, _ = import_draft(tmp_path, wording=Wording.OLD)

    # Both are 第三十條, and neither replaced the other.
    new_delay = next(c for c in store.clauses("享樂遊", Wording.NEW) if c.number == 30)
    old_delay = next(c for c in store.clauses("享樂遊", Wording.OLD) if c.number == 30)
    assert "次一班替代班機" in old_delay.text
    assert "次一班替代班機" not in new_delay.text
    assert store.products() == ["享樂遊"]

    book = openpyxl.load_workbook(old)
    [condition] = table(book["Conditions"], header_row=4)
    assert condition["Wording version"] == "old"
    assert condition["Delay-period rule"] == "old-wording rule"
    assert [(e["Exclusion type"], e["Clause reference"]) for e in table(book["Exclusions"])] == [
        ("other", "第四條 二"),
        ("strike at purchase", "第三十一條 二"),
        ("first replacement not taken", "第三十一條 四"),
    ]
    for workbook in (new, old):
        confirm(workbook)
        assert load_workbook(workbook, store=store).problems == ()


def imported(directory: Path, models: ScriptedModels) -> openpyxl.Workbook:
    """Import 享樂遊 in the new wording, and open the draft workbook."""
    workbook, _ = import_draft(directory, models)
    return openpyxl.load_workbook(workbook)
