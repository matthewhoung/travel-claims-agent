"""Judge a Scenario: a Verdict matrix with one cell per confirmed workbook.

Each test imports 享樂遊 in the new wording with scripted extraction,
confirms the workbook, scripts the facts read from the Scenario and the
provision judgements, and judges.
"""

from dataclasses import replace
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pytest

from travel_claims.app import CannotJudge, import_product, judge_scenario
from travel_claims.conditions import Benefit, ClauseRef
from travel_claims.judging import (
    Cell,
    ConditionOutcome,
    OutOfContract,
    ReadingOutcome,
    Reason,
    TurningProvision,
    Verdict,
)
from travel_claims.local_models import (
    ArrangedBy,
    Incident,
    Judgement,
    Leg,
    NeedsFact,
    NotSettled,
    ProvisionKind,
    Reading,
    Replacement,
    ScenarioFacts,
    Settled,
    TurnsOnCause,
)
from travel_claims.policies import Wording

from hsiang_le_you import (
    EXTRACTIONS,
    FIXTURE,
    FLIGHT_DELAY,
    FORCE_MAJEURE,
    NOTHING_APPLIES,
    POLICY_PERIOD,
    PROVISO,
    ROAD_CLOSURE,
    TAIPEI,
    confirm,
    import_draft,
    review,
    road_closure,
)
from scripted_models import ScriptedModels

SCENARIO = "我七月十日從桃園搭長榮班機去東京，班機延誤。"
SCHEDULED = datetime(2026, 7, 10, 8, 0, tzinfo=TAIPEI)


def delayed_by(delay: timedelta) -> ScenarioFacts:
    """The outbound flight from 桃園 departs `delay` late, within the policy period."""
    return ScenarioFacts(
        benefits=(Benefit.FLIGHT_DELAY,),
        incidents=(
            Incident(
                leg=Leg.OUTBOUND,
                airport="桃園國際機場",
                transport="flight",
                scheduled_departure=SCHEDULED,
                actual_departure=SCHEDULED + delay,
            ),
        ),
        cause="機械故障",
        purchased_at=datetime(2026, 7, 1, 12, 0, tzinfo=TAIPEI),
        policy_period=POLICY_PERIOD,
    )


def test_a_delay_of_exactly_four_hours_is_paid_with_one_step(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    models = ScriptedModels(facts=delayed_by(timedelta(hours=4)), judgements=NOTHING_APPLIES)

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert matrix.facts == delayed_by(timedelta(hours=4))
    assert matrix.cells == (
        Cell(
            product="享樂遊",
            wording=Wording.NEW,
            verdict=Verdict.PAID,
            outcomes=(
                ConditionOutcome(
                    condition="flight delay",
                    incident=1,
                    verdict=Verdict.PAID,
                    grounds="a delay of 4 h 0 min: 1 full step of 4 hours",
                    clause=ClauseRef(30),
                    delay=timedelta(hours=4),
                    steps=1,
                ),
            ),
        ),
    )


def test_a_delay_one_minute_short_of_four_hours_is_not_paid(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    # No judgement is scripted: a Condition whose threshold is not met is not judged further.
    models = ScriptedModels(facts=delayed_by(timedelta(hours=3, minutes=59)))

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert [cell.verdict for cell in matrix.cells] == [Verdict.NOT_PAID]
    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.NOT_PAID,
            grounds="a delay of 3 h 59 min, under the threshold of 4 hours",
            clause=ClauseRef(30),
            delay=timedelta(hours=3, minutes=59),
        ),
    )


def test_a_progressive_benefit_pays_each_full_step(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)

    def steps(delay: timedelta) -> int | None:
        models = ScriptedModels(facts=delayed_by(delay), judgements=NOTHING_APPLIES)
        matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)
        return matrix.cells[0].outcomes[0].steps

    assert steps(timedelta(hours=7, minutes=59)) == 1
    assert steps(timedelta(hours=8)) == 2


def test_a_one_off_benefit_pays_once_however_long_the_delay(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    review(workbook, {"Conditions!J5": "one-off fixed amount", "Conditions!K5": None})
    confirm(workbook)
    models = ScriptedModels(facts=delayed_by(timedelta(hours=8)), judgements=NOTHING_APPLIES)

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.PAID,
            grounds="a delay of 8 h 0 min: one payment",
            clause=ClauseRef(30),
            delay=timedelta(hours=8),
        ),
    )


def test_the_delay_runs_to_the_booked_flights_departure_not_a_later_flight_taken(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    booked = delayed_by(timedelta(hours=2, minutes=30))
    next_day = Replacement(
        departure=SCHEDULED + timedelta(days=1),
        arranged_by=ArrangedBy.INSURED,
        taken=True,
        arranged_at=SCHEDULED + timedelta(hours=1),
        destination="東京",
        returns_to_taiwan=False,
    )
    facts = replace(booked, incidents=(replace(booked.incidents[0], replacements=(next_day,)),))

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=ScriptedModels(facts=facts))

    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.NOT_PAID,
            grounds="a delay of 2 h 30 min, under the threshold of 4 hours",
            clause=ClauseRef(30),
            delay=timedelta(hours=2, minutes=30),
        ),
    )


def cancelled_return_flight(replacement_to: str, *, returns_to_taiwan: bool) -> ScenarioFacts:
    """The return flight from 成田 is cancelled at the end of the policy period.

    The insured arranges a replacement the next morning, after the policy expired.
    """
    scheduled = datetime(2026, 7, 14, 20, 0, tzinfo=TAIPEI)
    return ScenarioFacts(
        benefits=(Benefit.FLIGHT_DELAY,),
        incidents=(
            Incident(
                leg=Leg.RETURN,
                airport="成田國際機場",
                transport="flight",
                scheduled_departure=scheduled,
                cancelled=True,
                replacements=(
                    Replacement(
                        departure=datetime(2026, 7, 15, 14, 0, tzinfo=TAIPEI),
                        arranged_by=ArrangedBy.INSURED,
                        taken=True,
                        arranged_at=datetime(2026, 7, 15, 9, 0, tzinfo=TAIPEI),
                        destination=replacement_to,
                        returns_to_taiwan=returns_to_taiwan,
                    ),
                ),
            ),
        ),
        cause="颱風",
        purchased_at=datetime(2026, 7, 1, 12, 0, tzinfo=TAIPEI),
        policy_period=POLICY_PERIOD,
    )


def test_a_self_arranged_replacement_home_counts_though_arranged_after_the_policy_expired(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    facts = cancelled_return_flight("桃園", returns_to_taiwan=True)
    models = ScriptedModels(facts=facts, judgements=NOTHING_APPLIES)

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.PAID,
            grounds="a delay of 18 h 0 min: 4 full steps of 4 hours",
            clause=ClauseRef(30),
            delay=timedelta(hours=18),
            steps=4,
        ),
    )


def test_a_self_arranged_replacement_elsewhere_after_the_policy_expired_does_not_count(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    facts = cancelled_return_flight("首爾", returns_to_taiwan=False)

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=ScriptedModels(facts=facts))

    assert matrix.cells[0].verdict is Verdict.NOT_PAID
    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.NOT_PAID,
            grounds=(
                "no replacement flight counts toward the delay period: the one the insured "
                "arranged after the policy period ended does not return to Taiwan"
            ),
            clause=ClauseRef(30),
        ),
    )


def test_a_delay_stated_without_times_is_used_as_stated(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    timed = delayed_by(timedelta(0))
    stated = replace(timed.incidents[0], actual_departure=None, stated_delay=timedelta(hours=5))
    models = ScriptedModels(facts=replace(timed, incidents=(stated,)), judgements=NOTHING_APPLIES)

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert matrix.cells[0].outcomes[0].verdict is Verdict.PAID
    assert matrix.cells[0].outcomes[0].delay == timedelta(hours=5)


def test_a_booked_flight_scheduled_before_the_policy_period_is_not_paid(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    facts = replace(
        delayed_by(timedelta(hours=6)),
        policy_period=(
            datetime(2026, 7, 10, 9, 0, tzinfo=TAIPEI),
            datetime(2026, 7, 14, 23, 59, tzinfo=TAIPEI),
        ),
    )

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=ScriptedModels(facts=facts))

    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.NOT_PAID,
            grounds=(
                "the booked flight's scheduled departure, 2026-07-10 08:00, is outside "
                "the policy period, 2026-07-10 09:00 – 2026-07-14 23:59"
            ),
            clause=ClauseRef(30),
            delay=timedelta(hours=6),
        ),
    )


def test_a_coverage_requirement_the_facts_fall_outside_is_not_paid(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    review(workbook, {"Conditions!F5": "as a passenger"})  # judged as its own provision
    confirm(workbook)
    models = ScriptedModels(
        facts=delayed_by(timedelta(hours=5)),
        judgements={"第三十條": Settled(met=True), "as a passenger": Settled(met=False)},
    )

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert matrix.cells[0].outcomes[0].verdict is Verdict.NOT_PAID
    assert matrix.cells[0].outcomes[0].grounds == "outside the coverage requirement: as a passenger"
    assert matrix.cells[0].outcomes[0].clause == ClauseRef(30)


def test_an_exclusion_that_applies_is_not_paid_citing_it(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    models = ScriptedModels(
        facts=delayed_by(timedelta(hours=5)),
        judgements=NOTHING_APPLIES | {"第三十一條 二": Settled(met=True)},
    )

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.NOT_PAID,
            grounds=(
                "the exclusion applies: "
                "要保人向本公司申請訂立保險契約時，中華民國政府氣象機構已發布海上颱風警報。"
            ),
            clause=ClauseRef(31, "二"),
            delay=timedelta(hours=5),
        ),
    )


def test_each_confirmed_workbook_gets_its_own_cell_in_the_order_given(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    other = tmp_path / "享暢行.new.xlsx"
    import_product(
        FIXTURE,
        product="享暢行",
        wording=Wording.NEW,
        workbook=other,
        store=store,
        models=ScriptedModels(extractions=EXTRACTIONS),
    )
    review(other, {"Conditions!H5": 6})  # Threshold (hours)
    confirm(other)
    models = ScriptedModels(facts=delayed_by(timedelta(hours=5)), judgements=NOTHING_APPLIES)

    matrix = judge_scenario(SCENARIO, [other, workbook], store=store, models=models)

    assert [(cell.product, cell.wording, cell.verdict) for cell in matrix.cells] == [
        ("享暢行", Wording.NEW, Verdict.NOT_PAID),
        ("享樂遊", Wording.NEW, Verdict.PAID),
    ]


def test_nothing_is_judged_while_any_workbook_has_a_problem(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    unconfirmed = tmp_path / "享暢行.new.xlsx"
    import_product(
        FIXTURE,
        product="享暢行",
        wording=Wording.NEW,
        workbook=unconfirmed,
        store=store,
        models=ScriptedModels(extractions=EXTRACTIONS),
    )
    # Facts are scripted, but nothing may be asked of the models.
    models = ScriptedModels(facts=delayed_by(timedelta(hours=5)))

    with pytest.raises(CannotJudge) as refused:
        judge_scenario(SCENARIO, [workbook, unconfirmed], store=store, models=models)

    assert {
        path: [str(problem) for problem in problems]
        for path, problems in refused.value.problems.items()
    } == {
        unconfirmed: [
            "Conditions!B1: Confirmed by is missing",
            "Conditions!B2: Confirmed on is missing",
        ]
    }
    assert models.scenarios == []


# Undetermined outcomes -------------------------------------------------------------

TYPHOON_WARNING = "whether a sea typhoon warning was in force when the policy was bought"


def test_an_exclusion_that_needs_an_unstated_fact_leaves_the_outcome_undetermined(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    # A typhoon delays the flight, but the Scenario does not say whether a
    # warning was in force when the policy was bought.
    facts = replace(delayed_by(timedelta(hours=6)), cause="颱風")
    models = ScriptedModels(
        facts=facts, judgements=NOTHING_APPLIES | {"第三十一條 二": NeedsFact(TYPHOON_WARNING)}
    )

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert matrix.cells[0].verdict is Verdict.UNDETERMINED
    assert matrix.cells[0].reason is Reason.MISSING_FACT
    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.UNDETERMINED,
            reason=Reason.MISSING_FACT,
            grounds=f"the exclusion needs a fact the Scenario does not state: {TYPHOON_WARNING}",
            clause=ClauseRef(31, "二"),
            delay=timedelta(hours=6),
        ),
    )


def test_a_fact_a_window_needs_that_is_not_stated_leaves_the_outcome_undetermined(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    # Nothing is judged: a missing fact the code needs comes first.
    facts = replace(delayed_by(timedelta(hours=2)), policy_period=None)

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=ScriptedModels(facts=facts))

    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.UNDETERMINED,
            reason=Reason.MISSING_FACT,
            grounds="the Scenario does not state the policy period",
            clause=ClauseRef(30),
            delay=timedelta(hours=2),
        ),
    )


def test_a_cancelled_flight_with_no_replacement_time_lacks_the_end_of_the_delay(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    timed = delayed_by(timedelta(0))
    cancelled = replace(timed.incidents[0], actual_departure=None, cancelled=True)

    matrix = judge_scenario(
        SCENARIO,
        [workbook],
        store=store,
        models=ScriptedModels(facts=replace(timed, incidents=(cancelled,))),
    )

    assert matrix.cells[0].outcomes[0].reason is Reason.MISSING_FACT
    assert matrix.cells[0].outcomes[0].grounds == (
        "the Scenario does not state when the replacement flight departed"
    )


def test_the_road_closure_scenario_turns_on_the_cause_and_shows_both_readings(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    # The first replacement was not taken, so the exclusion applies unless
    # its force-majeure proviso does, which turns on the Cause.
    models = ScriptedModels(
        facts=road_closure(),
        judgements=NOTHING_APPLIES | {"第三十一條 五": Settled(met=True), PROVISO: FORCE_MAJEURE},
    )

    matrix = judge_scenario(ROAD_CLOSURE, [workbook], store=store, models=models)

    assert matrix.cells[0].verdict is Verdict.UNDETERMINED
    assert matrix.cells[0].reason is Reason.CAUSE_AMBIGUOUS
    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.UNDETERMINED,
            reason=Reason.CAUSE_AMBIGUOUS,
            grounds=f"turns on how the Cause is classified, under the proviso: {PROVISO}",
            clause=ClauseRef(31, "五"),
            delay=timedelta(hours=5),
            turns_on=(
                TurningProvision(
                    kind=ProvisionKind.PROVISO,
                    text=PROVISO,
                    clause=ClauseRef(31, "五"),
                    readings=(
                        ReadingOutcome(
                            cause="不可抗力",
                            verdict=Verdict.PAID,
                            grounds="a delay of 5 h 0 min: 1 full step of 4 hours",
                            clause=ClauseRef(30),
                            delay=timedelta(hours=5),
                        ),
                        ReadingOutcome(
                            cause="非不可抗力",
                            verdict=Verdict.NOT_PAID,
                            grounds=(
                                "the exclusion applies: "
                                "被保險人未搭乘航空業者所提供之第一班替代交通工具。"
                            ),
                            clause=ClauseRef(31, "五"),
                            delay=timedelta(hours=5),
                        ),
                    ),
                ),
            ),
        ),
    )


def test_provisions_that_turn_on_the_cause_are_read_together_by_classification(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    # The wilful-act exclusion and the force-majeure proviso both turn on
    # whether the road closure is force majeure. Under each reading both are
    # settled, so each reading has a definite outcome.
    models = ScriptedModels(
        facts=road_closure(),
        judgements=NOTHING_APPLIES
        | {
            "第四條 二": TurnsOnCause(
                (Reading(cause="不可抗力", met=False), Reading(cause="非不可抗力", met=True))
            ),
            "第三十一條 五": Settled(met=True),
            PROVISO: FORCE_MAJEURE,
        },
    )

    matrix = judge_scenario(ROAD_CLOSURE, [workbook], store=store, models=models)

    outcome = matrix.cells[0].outcomes[0]
    assert outcome.reason is Reason.CAUSE_AMBIGUOUS
    assert [
        (str(p.clause), p.kind, [(r.cause, r.verdict, str(r.clause)) for r in p.readings])
        for p in outcome.turns_on
    ] == [
        (
            "第四條 二",
            ProvisionKind.EXCLUSION,
            [("不可抗力", Verdict.PAID, "第三十條"), ("非不可抗力", Verdict.NOT_PAID, "第四條 二")],
        ),
        (
            "第三十一條 五",
            ProvisionKind.PROVISO,
            [("不可抗力", Verdict.PAID, "第三十條"), ("非不可抗力", Verdict.NOT_PAID, "第四條 二")],
        ),
    ]


def test_a_scenario_that_states_the_trip_is_outside_the_policy_period_is_not_paid(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    facts = replace(road_closure(), within_policy_period=False)

    matrix = judge_scenario(
        ROAD_CLOSURE, [workbook], store=store, models=ScriptedModels(facts=facts)
    )

    assert matrix.cells[0].outcomes[0].verdict is Verdict.NOT_PAID
    assert matrix.cells[0].outcomes[0].grounds == (
        "the Scenario states the trip is outside the policy period"
    )


def test_a_proviso_that_applies_lifts_its_exclusion(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)

    def verdict(proviso: Judgement) -> Verdict:
        models = ScriptedModels(
            facts=road_closure(),
            judgements=NOTHING_APPLIES | {"第三十一條 五": Settled(met=True), PROVISO: proviso},
        )
        matrix = judge_scenario(ROAD_CLOSURE, [workbook], store=store, models=models)
        return matrix.cells[0].verdict

    assert verdict(Settled(met=True)) is Verdict.PAID
    assert verdict(Settled(met=False)) is Verdict.NOT_PAID


def test_a_proviso_is_not_judged_when_its_exclusion_does_not_apply(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    # NOTHING_APPLIES scripts no answer for the proviso; asking for one fails the test.
    models = ScriptedModels(facts=road_closure(), judgements=NOTHING_APPLIES)

    matrix = judge_scenario(ROAD_CLOSURE, [workbook], store=store, models=models)

    assert matrix.cells[0].verdict is Verdict.PAID


def test_clauses_that_do_not_settle_the_situation_leave_the_outcome_undetermined(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    unaddressed = "a flight that took off on time and returned to the airport it left"
    models = ScriptedModels(
        facts=delayed_by(timedelta(hours=5)),
        judgements=NOTHING_APPLIES | {FLIGHT_DELAY.covered_event: NotSettled(unaddressed)},
    )

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert matrix.cells[0].reason is Reason.CLAUSES_SILENT
    assert matrix.cells[0].outcomes == (
        ConditionOutcome(
            condition="flight delay",
            incident=1,
            verdict=Verdict.UNDETERMINED,
            reason=Reason.CLAUSES_SILENT,
            grounds=f"the Clauses do not settle the covered event: {unaddressed}",
            clause=ClauseRef(30),
            delay=timedelta(hours=5),
        ),
    )


def test_an_exclusion_that_applies_beats_one_that_turns_on_the_cause(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    # 第四條 二, the wilful-act exclusion, concerns the Cause.
    models = ScriptedModels(
        facts=delayed_by(timedelta(hours=5)),
        judgements=NOTHING_APPLIES
        | {
            "第四條 二": TurnsOnCause((Reading("故意", met=True), Reading("過失", met=False))),
            "第三十一條 二": Settled(met=True),
        },
    )

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert matrix.cells[0].verdict is Verdict.NOT_PAID
    assert matrix.cells[0].outcomes[0].clause == ClauseRef(31, "二")
    assert matrix.cells[0].outcomes[0].reason is None


def test_a_provision_that_does_not_concern_the_cause_may_not_turn_on_it(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    # The typhoon-warning exclusion is about what was in force at purchase, not the Cause.
    models = ScriptedModels(
        facts=delayed_by(timedelta(hours=5)),
        judgements=NOTHING_APPLIES
        | {
            "第三十一條 二": TurnsOnCause(
                (Reading("颱風", met=True), Reading("機械故障", met=False))
            )
        },
    )

    with pytest.raises(OutOfContract) as broken:
        judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert str(broken.value) == (
        "the local models answered that the exclusion of 第三十一條 二 turns on the Cause, "
        "which only a provision that concerns the Cause may: "
        "要保人向本公司申請訂立保險契約時，中華民國政府氣象機構已發布海上颱風警報。"
    )


def test_an_undetermined_cell_takes_its_reason_by_precedence_and_each_condition_keeps_its_own(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    # The reviewer splits the flight-delay Condition in two, by covered event.
    book = openpyxl.load_workbook(workbook)
    sheet = book["Conditions"]
    for column in range(1, sheet.max_column + 1):
        sheet.cell(6, column, sheet.cell(5, column).value)
    book.save(workbook)
    review(
        workbook,
        {
            "Conditions!A5": "flight delay / delayed",
            "Conditions!E5": "the scheduled flight departs 4 hours or more late",
            "Conditions!A6": "flight delay / cancelled",
            "Conditions!E6": "the scheduled flight is cancelled",
        },
    )
    confirm(workbook)
    unaddressed = "a flight that took off on time and returned to the airport it left"
    models = ScriptedModels(
        facts=delayed_by(timedelta(hours=5)),
        judgements=NOTHING_APPLIES
        | {
            "the scheduled flight departs 4 hours or more late": NotSettled(unaddressed),
            "the scheduled flight is cancelled": NeedsFact("whether the flight was cancelled"),
        },
    )

    matrix = judge_scenario(SCENARIO, [workbook], store=store, models=models)

    assert (matrix.cells[0].verdict, matrix.cells[0].reason) == (
        Verdict.UNDETERMINED,
        Reason.MISSING_FACT,
    )
    assert [(o.condition, o.verdict, o.reason) for o in matrix.cells[0].outcomes] == [
        ("flight delay / delayed", Verdict.UNDETERMINED, Reason.CLAUSES_SILENT),
        ("flight delay / cancelled", Verdict.UNDETERMINED, Reason.MISSING_FACT),
    ]
