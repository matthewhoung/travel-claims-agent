"""Load a workbook: validate what the reviewer confirmed, and list every problem at once.

Each test imports 享樂遊 with scripted extraction, edits the draft workbook
as a reviewer would in Excel, and loads it.
"""

from datetime import date
from pathlib import Path

import openpyxl

from travel_claims.app import import_product, load_workbook
from travel_claims.clause_store import ClauseStore
from travel_claims.conditions import (
    Benefit,
    BenefitType,
    ClauseRef,
    Condition,
    DelayPeriodRule,
    Exclusion,
    ExclusionType,
)
from travel_claims.loading import ConditionTable
from travel_claims.policies import Wording

from hsiang_le_you import FIXTURE, confirm, import_draft, review
from scripted_models import ScriptedModels


def test_a_draft_nobody_has_confirmed_lacks_the_confirmation_record(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)

    loaded = load_workbook(workbook, store=store)

    assert [str(problem) for problem in loaded.problems] == [
        "Conditions!B1: Confirmed by is missing",
        "Conditions!B2: Confirmed on is missing",
    ]
    assert loaded.table is None


def test_a_confirmed_workbook_loads_as_the_condition_table(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)

    loaded = load_workbook(workbook, store=store)

    assert loaded.problems == ()
    assert loaded.table == ConditionTable(
        product="享樂遊",
        wording=Wording.NEW,
        confirmed_by="王小明",
        confirmed_on=date(2026, 10, 1),
        conditions=(
            Condition(
                key="flight delay",
                product="享樂遊",
                wording=Wording.NEW,
                benefit=Benefit.FLIGHT_DELAY,
                covered_event="scheduled flight departs 4 hours or more late",
                coverage_requirements=("scheduled flight", "as a passenger"),
                coverage_window="the dates and times on the policy schedule",
                threshold_hours=4,
                delay_period_rule=DelayPeriodRule.NEW,
                benefit_type=BenefitType.PROGRESSIVE,
                step_hours=4,
                max_claims_per_period=2,
                aggregate_limit_group=None,
                eligible_costs=(),
                cost_maximums=None,
                clause=ClauseRef(30),
            ),
        ),
        exclusions=(
            Exclusion(
                product="享樂遊",
                wording=Wording.NEW,
                type=ExclusionType.OTHER,
                applies_to=("all",),
                text="被保險人故意行為。",
                proviso=None,
                concerns_cause=True,
                clause=ClauseRef(4, "二"),
            ),
            Exclusion(
                product="享樂遊",
                wording=Wording.NEW,
                type=ExclusionType.TYPHOON_WARNING,
                applies_to=("flight delay",),
                text="要保人向本公司申請訂立保險契約時，中華民國政府氣象機構已發布海上颱風警報。",
                proviso=None,
                concerns_cause=False,
                clause=ClauseRef(31, "二"),
            ),
            Exclusion(
                product="享樂遊",
                wording=Wording.NEW,
                type=ExclusionType.FIRST_REPLACEMENT_NOT_TAKEN,
                applies_to=("flight delay",),
                text="被保險人未搭乘航空業者所提供之第一班替代交通工具。",
                proviso="但被保險人因不可抗力因素致無法搭乘航空業者所提供之第一班替代交通工具者，"
                "不在此限。",
                concerns_cause=True,
                clause=ClauseRef(31, "五"),
            ),
        ),
    )
    # 14 fields of the Condition and 6 of each exclusion; Product and Wording
    # version were chosen at import, not extracted.
    assert (loaded.extracted_fields, loaded.changed_fields) == (32, 0)


def test_every_problem_in_the_workbook_is_listed_at_once(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    review(
        workbook,
        {
            # Conditions: the flight-delay Condition is on row 5.
            "Conditions!H5": None,  # Threshold (hours)
            "Conditions!J5": "progresive",  # Benefit type
            # Exclusions: the wilful act, the typhoon warning and the first
            # replacement flight are on rows 2 to 4.
            "Exclusions!C2": None,  # Exclusion type
            "Exclusions!D3": "flight delays",  # Applies to
            "Exclusions!H4": "第三十九條 五",  # Clause reference
            # Amounts: a row for a Condition the workbook does not have.
            "Amounts!A2": "flight delays",
            "Amounts!C2": "not collected",
        },
    )

    loaded = load_workbook(workbook, store=store)

    assert [str(problem) for problem in loaded.problems] == [
        "Conditions!B1: Confirmed by is missing",
        "Conditions!B2: Confirmed on is missing",
        "Conditions!H5: Threshold (hours) is missing",
        (
            "Conditions!J5: Benefit type is not one of: "
            "progressive fixed amount, one-off fixed amount, reimbursement"
        ),
        "Exclusions!C2: Exclusion type is missing",
        "Exclusions!D3: unknown Condition key: flight delays",
        "Exclusions!H4: 第三十九條 is not among the stored Clauses of 享樂遊 (new wording)",
        "Amounts!A2: unknown Condition key: flight delays",
    ]
    assert loaded.table is None


def test_load_counts_the_extracted_fields_the_reviewer_changed(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    review(
        workbook,
        {
            "Conditions!E5": "scheduled flight departs at least 4 hours late",  # Covered event
            "Conditions!H5": "4",  # Threshold (hours), typed again as text: not a change
            "Conditions!M5": "flight delay",  # Aggregate limit group, left empty by extraction
            "Exclusions!G3": "yes",  # Concerns Cause of the typhoon warning
        },
    )
    book = openpyxl.load_workbook(workbook)
    exclusions = book["Exclusions"]
    exclusions.delete_rows(2)  # the wilful act
    exclusions.append(
        ["享樂遊", "new", "other", "all", "被保險人犯罪行為。", None, "yes", "第四條 一"]
    )
    book.save(workbook)

    loaded = load_workbook(workbook, store=store)

    assert loaded.problems == ()
    # Three fields edited, and the six fields of the deleted row. The added
    # row was not extracted, so none of its fields count.
    assert (loaded.extracted_fields, loaded.changed_fields) == (32, 9)


def test_clause_references_are_checked_against_the_stored_clauses_of_the_wording_version(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    review(workbook, {"Exclusions!H2": "第十八條"})  # stored at import, though not extracted
    old_only = ClauseStore(tmp_path / "old-only")
    import_product(
        FIXTURE,
        product="享樂遊",
        wording=Wording.OLD,
        workbook=tmp_path / "享樂遊.old.xlsx",
        store=old_only,
        models=ScriptedModels(),
    )

    assert load_workbook(workbook, store=store).problems == ()
    assert [str(p) for p in load_workbook(workbook, store=old_only).problems] == [
        "Conditions!P5: 第三十條 is not among the stored Clauses of 享樂遊 (new wording)",
        "Exclusions!H2: 第十八條 is not among the stored Clauses of 享樂遊 (new wording)",
        "Exclusions!H3: 第三十一條 is not among the stored Clauses of 享樂遊 (new wording)",
        "Exclusions!H4: 第三十一條 is not among the stored Clauses of 享樂遊 (new wording)",
    ]


def test_an_ambiguous_condition_key_or_a_row_of_another_product_is_a_problem(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    book = openpyxl.load_workbook(workbook)
    conditions = book["Conditions"]
    conditions.append([cell.value for cell in conditions[5]])
    book["Exclusions"]["A3"] = "享暢行"
    book.save(workbook)

    loaded = load_workbook(workbook, store=store)

    # With two Conditions, the key "flight delay" would also name both of them.
    benefit_name = "Condition key flight delay is also the name of its Benefit, which has several"
    assert [str(problem) for problem in loaded.problems] == [
        f"Conditions!A5: {benefit_name} Conditions",
        "Conditions!A6: Condition key flight delay is also in row 5",
        f"Conditions!A6: {benefit_name} Conditions",
        "Exclusions!A3: Product is 享暢行, but the workbook is for 享樂遊 (new wording)",
    ]


def test_rows_sorted_in_excel_or_with_a_corrected_key_are_still_matched(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    book = openpyxl.load_workbook(workbook)
    exclusions = book["Exclusions"]
    rows = [[cell.value for cell in row] for row in exclusions.iter_rows(min_row=2)]
    for number, cells in enumerate(reversed(rows), start=2):
        for column, value in enumerate(cells, start=1):
            exclusions.cell(number, column).value = value
    exclusions["H2"] = "第三十一條 四"  # the first-replacement exclusion, now on row 2
    book.save(workbook)

    loaded = load_workbook(workbook, store=store)

    assert loaded.problems == ()
    assert (loaded.extracted_fields, loaded.changed_fields) == (32, 1)


def test_an_exclusion_type_must_be_on_the_list_of_every_benefit_it_applies_to(
    tmp_path: Path,
) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    review(workbook, {"Exclusions!D4": "flight delay; trip change"})  # first replacement not taken

    loaded = load_workbook(workbook, store=store)

    assert [str(problem) for problem in loaded.problems] == [
        (
            "Exclusions!C4: Exclusion type first replacement not taken"
            " is not on the list for trip change"
        )
    ]


def test_amounts_rows_may_name_a_condition_or_all(tmp_path: Path) -> None:
    workbook, store = import_draft(tmp_path)
    confirm(workbook)
    review(
        workbook,
        {
            "Amounts!A2": "flight delay",
            "Amounts!C2": "not published",
            "Amounts!A3": "all",
            "Amounts!C3": "not collected",
        },
    )

    assert load_workbook(workbook, store=store).problems == ()
